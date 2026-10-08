"""Processor for PowerPoint presentations (.pptx).

Extracts per-slide text (titles as ``# Title`` headings, body text frames,
and tables). Slide boundaries are published as ``meta.segments`` per the IR
contract. ``pages`` (alias ``slides``) picks slides; ``images: true`` adds
a picture of every slide, drawn by LibreOffice (``_office_pages``), with
the slide's number as ``page``; ``embedded_images: true`` adds the
pictures stored in the slides.

Requires ``python-pptx``: ``pip install attachments[pptx]``.

python-pptx only reads ``.pptx``; ``.ppt`` and ``.odp`` are converted by
LibreOffice first (``legacy_office.py``).
"""

from __future__ import annotations

import io
import logging
from typing import Any

from .._options import Option, register_options
from .._pages import PageSelectionError, describe_pages, picker_from_options
from ..types import (
    ERROR_INVALID_OPTION,
    ERROR_PARSE,
    error_artifact,
    make_artifact,
    missing_dep_artifact,
)
from . import register_processor
from ._office_pages import (
    DEFAULT_DPI,
    DEFAULT_MAX_DIM,
    check_render_options,
    office_pictures,
    render_options,
    wants_pictures,
)

log = logging.getLogger("attachments.processors.pptx")

_SLIDE_SEPARATOR = "\n\n"


def _join_slides_with_segments(
    slides: list[tuple[str, str]],
    numbers: list[int] | None = None,
) -> tuple[str, list[dict]]:
    """Join per-slide texts and build slide segments (IR contract: meta.segments).

    *slides* is a list of ``(label, slide_text)`` pairs, one per slide in
    deck order; *numbers* are their 1-based slide numbers (default 1, 2, ...;
    a selection of slides keeps the deck's numbers). Each segment carries
    the slide label (its title, which may
    be any text), the 1-based slide number as ``page`` (the same number as
    ``ImageItem.page`` on that slide's pictures), and start/end offsets
    that slice the joined text exactly back to that slide's text.

    Examples:
        >>> text, segs = _join_slides_with_segments(
        ...     [("Intro", "# Intro\\nhello"), ("slide 2", "world")]
        ... )
        >>> text
        '# Intro\\nhello\\n\\nworld'
        >>> segs[0]
        {'kind': 'slide', 'label': 'Intro', 'start': 0, 'end': 13, 'page': 1}
        >>> text[segs[1]["start"] : segs[1]["end"]]
        'world'
    """
    parts: list[str] = []
    segments: list[dict] = []
    offset = 0
    numbers = numbers or list(range(1, len(slides) + 1))
    for slide_no, (label, slide_text) in zip(numbers, slides, strict=True):
        if parts:
            offset += len(_SLIDE_SEPARATOR)
        segments.append(
            {
                "kind": "slide",
                "label": label,
                "start": offset,
                "end": offset + len(slide_text),
                "page": slide_no,
            }
        )
        parts.append(slide_text)
        offset += len(slide_text)
    return _SLIDE_SEPARATOR.join(parts), segments


def _walk_shapes(shapes):
    """Yield shapes depth-first, descending into group shapes.

    Group shapes expose a nested ``.shapes`` collection; everything else
    is yielded as-is (duck typing keeps this import-free).
    """
    for shape in shapes:
        nested = getattr(shape, "shapes", None)
        if nested is not None:
            yield from _walk_shapes(nested)
        else:
            yield shape


def _table_text(table) -> str:
    """Render a pptx table as a simple pipe-delimited text table."""
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    return "\n".join(rows)


def _slide_text(slide) -> tuple[str, str | None]:
    """Return ``(slide_text, title)`` for one slide.

    The title (when present and non-empty) becomes a ``# <title>``
    heading line; remaining text frames and tables follow in shape
    order, joined with newlines.
    """
    try:
        title_shape = slide.shapes.title
    except Exception:  # malformed layouts must never break extraction
        title_shape = None

    title: str | None = None
    if title_shape is not None and getattr(title_shape, "has_text_frame", False):
        title = title_shape.text_frame.text.strip() or None

    parts: list[str] = []
    if title:
        parts.append(f"# {title}")
    for shape in _walk_shapes(slide.shapes):
        # BaseShape.__eq__ compares the underlying XML element ("is" would
        # fail: python-pptx builds fresh shape wrappers on every iteration).
        if title_shape is not None and shape == title_shape:
            continue  # already emitted as the heading line
        if getattr(shape, "has_table", False):
            parts.append(_table_text(shape.table))
        elif getattr(shape, "has_text_frame", False):
            text = shape.text_frame.text.strip()
            if text:
                parts.append(text)
    return "\n".join(parts), title


def pptx_processor(
    data: bytes,
    *,
    filename: str | None = None,
    render_images: bool | str = False,
    embedded_images: bool = False,
    dpi: int = DEFAULT_DPI,
    max_dim: int | None = DEFAULT_MAX_DIM,
    image_format: str = "png",
    quality: int | None = None,
    _render_from: tuple[bytes, str] | None = None,
    **_opts: Any,
) -> dict[str, Any]:
    """PowerPoint (.pptx) -> artifact dict (text, images, audio, video, meta).

    Options (via att(..., **options)):
      - pages / slides (``page_start``/``page_end``/``page_selection``):
        which slides, e.g. ``2-5``, ``1,3``, ``-1`` (the last slide).
      - render_images (DSL ``images``): ``True`` adds a picture of every
        (selected) slide, drawn by LibreOffice, with its slide number as
        ``page``; ``"auto"`` only when LibreOffice is installed.
      - embedded_images: the pictures stored in the (selected) slides,
        each with its slide number as ``page``.
      - dpi / max_dim / image_format / quality: slide picture size/format.
      - _render_from: ``(bytes, ext)`` to draw the slides from instead of
        *data* (``.ppt``/``.odp`` are drawn from the original file).

    meta.kind is ``"slides"``; meta.segments has one ``"slide"`` segment
    per slide whose offsets slice ``text`` exactly to that slide's text.
    """
    source = filename or "file.pptx"
    try:
        pick = picker_from_options(_opts)
    except PageSelectionError as e:
        artifact = error_artifact(source, ERROR_INVALID_OPTION, str(e))
        artifact["meta"]["kind"] = "slides"
        return artifact
    mode = wants_pictures(render_images)
    if mode:
        invalid = check_render_options(
            source=source,
            kind="slides",
            dpi=dpi,
            max_dim=max_dim,
            image_format=image_format,
            quality=quality,
        )
        if invalid:
            return invalid

    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError:
        return missing_dep_artifact(source, "pptx")

    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as e:
        log.warning("failed to open pptx %s: %s", source, e)
        artifact = error_artifact(
            source, ERROR_PARSE, f"Failed to parse PowerPoint file: {e}"
        )
        artifact["meta"]["kind"] = "slides"
        return artifact

    try:
        # --- per-slide text + segments ---
        all_slides = list(prs.slides)
        chosen = pick(len(all_slides)) if pick is not None else range(len(all_slides))
        numbers = [i + 1 for i in chosen]
        entries: list[tuple[str, str]] = []
        for slide_no in numbers:
            slide_text, title = _slide_text(all_slides[slide_no - 1])
            entries.append((title or f"slide {slide_no}", slide_text))
        text, segments = _join_slides_with_segments(entries, numbers)

        extra: dict[str, Any] = {"slide_count": len(all_slides)}
        if pick is not None:
            extra["slides_selected"] = numbers

        # --- embedded picture extraction (optional) ---
        images: list[dict[str, Any]] = []
        if embedded_images:
            stem = filename or "presentation"
            for slide_no in numbers:
                slide = all_slides[slide_no - 1]
                for shape in _walk_shapes(slide.shapes):
                    try:
                        if shape.shape_type != MSO_SHAPE_TYPE.PICTURE and not hasattr(
                            shape, "image"
                        ):
                            continue
                        # Raises for linked (not embedded) pictures.
                        image = shape.image
                        mimetype = image.content_type or "image/png"
                        ext = image.ext or mimetype.split("/")[-1].split("+")[0]
                        images.append(
                            {
                                "name": (
                                    f"{stem}-slide-{slide_no}-"
                                    f"image-{len(images) + 1}.{ext}"
                                ),
                                "mimetype": mimetype,
                                "bytes": image.blob,
                                "page": slide_no,
                            }
                        )
                    except Exception as exc:
                        log.debug("skipping picture on slide %d: %s", slide_no, exc)
            extra["images_extracted"] = len(images)
    except Exception as e:
        log.warning("failed to extract pptx content from %s: %s", source, e)
        artifact = error_artifact(
            source, ERROR_PARSE, f"Failed to parse PowerPoint file: {e}"
        )
        artifact["meta"]["kind"] = "slides"
        return artifact

    meta: dict[str, Any] = {"kind": "slides", "extra": extra}
    if segments:
        meta["segments"] = segments
    if pick is not None and not numbers:
        meta["warnings"] = [
            f"pages: {describe_pages(_opts)} selects no slide of this "
            f"{len(all_slides)}-slide deck"
        ]
    elif mode:
        pictures = office_pictures(
            *(_render_from or (data, ".pptx")),
            source=source,
            kind="slides",
            mode=mode,
            pick=pick,
            dpi=dpi,
            max_dim=max_dim,
            image_format=image_format,
            quality=quality or 85,
        )
        if isinstance(pictures, dict):  # an error artifact
            return pictures
        slide_images, picture_extra, note = pictures
        # Each slide's picture first, then the pictures embedded in it.
        images = sorted(slide_images + images, key=lambda im: im["page"])
        extra.update(picture_extra)
        if note:
            meta["note"] = note
    return make_artifact(text=text, images=images, meta=meta)


OPTIONS: tuple[Option, ...] = (
    Option(
        "pages",
        "pages",
        aliases=("slides", "page"),
        help="Slides to read: 3, 2-5, 7- (to the end), 1,3,5, -1 (last)",
        example="pages: 1-5",
    ),
    *render_options(unit="slide"),
)


def register() -> None:
    """Register the pptx processor and its option schema (idempotent)."""
    register_processor(".pptx", pptx_processor)
    register_options(".pptx", OPTIONS)


register()
