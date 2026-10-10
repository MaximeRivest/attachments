"""The 20 tasks of evals/bench: files, prompts, graders, expert and naive solutions.

Five come from evals/skill (photos, charts, folder, install, scans) unchanged;
fifteen are new: attachments with lm15, real model calls, MCP and the CLI.
Every task has a trap seen in real use; the grader checks outcomes only.
"""

from __future__ import annotations

import io
import json
import random
import subprocess
import sys
import textwrap
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "skill"))
import run as skill_eval  # noqa: E402  (evals/skill/run.py)

#: The model every task that calls a model must use (through lm15).
TASK_MODEL = "openai-codex:gpt-6-luna"

LIBS = (
    "Use the Python libraries attachments and lm15, already installed in .venv "
    "(run Python as .venv/bin/python)."
)
CALL = (
    f"To ask the model, use lm15: its router reaches `{TASK_MODEL}` by that name "
    "(this machine is already signed in)."
)
MCP = "You have the attachments tools (MCP) for reading files, and no shell. Use them."


@dataclass
class Task:
    prompt: str
    build: Callable[[Path], None]
    grader: str | None  # None: the evals/skill grader
    expert: dict = field(default_factory=dict)
    naive: dict = field(default_factory=dict)
    venv: bool = True
    naive_may_pass: str = ""  # why the naive solution no longer fails
    mcp: bool = False
    calls_model: bool = False
    site: bool = False  # prompt has {site}: the local web server


# =============================================================================
# File builders
# =============================================================================


def _font(size: int):
    from PIL import ImageFont

    return ImageFont.load_default(size=size)


def _text_pdf(path: Path, pages: list[str], size=(612, 792), fontsize=11) -> None:
    import pymupdf

    doc = pymupdf.open()
    for body in pages:
        page = doc.new_page(width=size[0], height=size[1])
        page.insert_textbox(
            pymupdf.Rect(54, 54, size[0] - 54, size[1] - 54), body, fontsize=fontsize
        )
    doc.save(path)
    doc.close()


def _scan_page(lines: list[str]) -> bytes:
    """A scanned page: text drawn into a JPEG, no text layer."""
    import numpy as np
    from PIL import Image, ImageDraw

    img = Image.new("L", (1275, 1650), 240)
    draw = ImageDraw.Draw(img)
    y = 120
    for line in lines:
        draw.text((110, y), line, fill=25, font=_font(34))
        y += 60
    rng = np.random.default_rng(len(lines))
    noisy = np.array(img, dtype=np.float32) + rng.normal(0, 6, (1650, 1275))
    img = Image.fromarray(noisy.clip(0, 255).astype("uint8")).rotate(0.3, fillcolor=240)
    buf = io.BytesIO()
    img.convert("RGB").save(buf, "JPEG", quality=85)
    return buf.getvalue()


def _scanned_pdf(path: Path, pages: list[list[str]]) -> None:
    import pymupdf

    doc = pymupdf.open()
    for lines in pages:
        page = doc.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=_scan_page(lines))
    doc.save(path)
    doc.close()


def build_sheet(work: Path) -> None:
    import openpyxl

    book = openpyxl.Workbook()
    summary = book.active
    summary.title = "Summary"
    summary.append(["note", "SUMMARY-ONLY: totals are computed elsewhere"])
    sales = book.create_sheet("Sales")
    sales.append(["order", "region", "amount"])
    rng = random.Random(6)
    for n in range(1, 5001):
        sales.append(
            [
                f"S-{n:05d}",
                rng.choice(["North", "South", "East", "West"]),
                rng.randint(10, 999),
            ]
        )
    costs = book.create_sheet("Costs")
    costs.append(["item", "cost"])
    costs.append(["COSTS-ONLY servers", 1200])
    book.save(work / "data.xlsx")


HANDBOOK = [
    "Employee handbook, part DOC-P1.\nVacation: 25 days a year, booked two weeks ahead.",
    "Employee handbook, part DOC-P2.\nRemote work: up to three days a week.",
    "Employee handbook, part DOC-P3.\nExpenses: receipts within 30 days.",
]


def build_handbook(work: Path) -> None:
    _text_pdf(work / "handbook.pdf", HANDBOOK)


def build_charts(work: Path) -> None:
    skill_eval.build_files("charts", work)


def build_deck(work: Path) -> None:
    """8 slides, the same title on each, a chart; no slide number anywhere."""
    import pymupdf

    rng = random.Random(9)
    doc = pymupdf.open()
    for _ in range(8):
        page = doc.new_page(width=960, height=540)
        page.insert_text((60, 70), "Regional sales", fontsize=26)
        color = (rng.random() * 0.8, rng.random() * 0.8, rng.random() * 0.8)
        for i in range(4):
            h = rng.randint(80, 330)
            x = 120 + i * 180
            page.draw_rect(
                pymupdf.Rect(x, 480 - h, x + 110, 480), color=color, fill=color
            )
    doc.save(work / "deck.pdf")
    doc.close()


WORDS = (
    "valve pump filter sensor gasket bracket housing seal motor relay coupling "
    "bearing manifold nozzle hose clamp switch panel cable fuse"
).split()


def manual_pages() -> list[str]:
    rng = random.Random(10)
    pages = []
    for n in range(1, 61):
        sentences = [f"Manual page {n} [P-{n:03d}]."]
        while sum(len(s) + 1 for s in sentences) < 1700:
            a, b, c = rng.sample(WORDS, 3)
            sentences.append(
                f"Check the {a} next to the {b} before replacing the {c} ({n}.{len(sentences)})."
            )
        pages.append(" ".join(sentences))
    return pages


def build_manual(work: Path) -> None:
    _text_pdf(work / "manual.pdf", manual_pages(), fontsize=9)


def build_project(work: Path) -> None:
    proj = work / "proj"
    (proj / "app").mkdir(parents=True)
    (proj / "app" / "server.py").write_text(
        "import os\n\nPORT = int(os.environ.get('PORT', 8080))  # CODE-SERVER\n\n"
        "def handle(request):\n    return {'status': 'ok'}\n"
    )
    (proj / "app" / "billing.py").write_text(
        "def invoice_total(lines):  # CODE-BILLING\n    return sum(q * p for q, p in lines)\n"
    )
    (proj / "README.md").write_text(
        "# Shop backend\nREADME-MARKER: run `python -m app.server`.\n"
    )
    (proj / "settings.toml").write_text("[server]\nworkers = 4  # SETTINGS-MARKER\n")
    (proj / ".env").write_text(
        "STRIPE_KEY=sk_live_SECRET_9f8e7d6c5b4a\nDB_PASSWORD=hunter2-SECRET\n"
    )
    (proj / "deploy_key").write_text(
        "-----BEGIN OPENSSH PRIVATE KEY-----\nb3BlbnNzaC1rZXktdjEAAAAA-SECRET-KEY\n"
        "-----END OPENSSH PRIVATE KEY-----\n"
    )


INVOICES = {
    "acme_tools.pdf": (1284.50, False),
    "brightline_print.pdf": (932.10, True),
    "northwind_freight.pdf": (4410.00, False),
    "office_supply_co.pdf": (77.25, False),
    "quarry_stone.pdf": (2650.80, True),
    "zenith_it.pdf": (15999.99, False),
}


def _invoice_lines(name: str, total: float) -> list[str]:
    vendor = name.removesuffix(".pdf").replace("_", " ").title()
    sub = round(total / 1.15, 2)
    return [
        f"INVOICE  {vendor}",
        f"Invoice no. {abs(hash(name)) % 90000 + 10000}",
        "Bill to: Rivest Labs",
        f"Services rendered ........ ${sub:,.2f}",
        f"Tax (15%) ................ ${total - sub:,.2f}",
        f"TOTAL DUE ................ ${total:,.2f}",
        "Payment within 30 days.",
    ]


def build_invoices(work: Path) -> None:
    folder = work / "invoices"
    folder.mkdir()
    for name, (total, scanned) in INVOICES.items():
        lines = _invoice_lines(name, total)
        if scanned:
            _scanned_pdf(folder / name, [lines])
        else:
            _text_pdf(folder / name, ["\n".join(lines)])


QUARTERS = {
    "Q1": {"North": 120, "South": 95, "East": 143, "West": 88},
    "Q2": {"North": 131, "South": 162, "East": 117, "West": 90},
    "Q3": {"North": 101, "South": 99, "East": 112, "West": 176},
    "Q4": {"North": 188, "South": 140, "East": 151, "West": 129},
}


def _chart_png(values: dict) -> bytes:
    from PIL import Image, ImageDraw

    img = Image.new("RGB", (1000, 560), "white")
    draw = ImageDraw.Draw(img)
    for i, (region, v) in enumerate(values.items()):
        x = 90 + i * 220
        top = 470 - v * 2
        draw.rectangle((x, top, x + 140, 470), fill=(40, 90, 170))
        draw.text((x + 40, top - 40), str(v), fill="black", font=_font(30))
        draw.text((x + 20, 485), region, fill="black", font=_font(30))
    buf = io.BytesIO()
    img.save(buf, "PNG")
    return buf.getvalue()


def build_sales(work: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    for quarter, values in QUARTERS.items():
        page = doc.new_page(width=842, height=595)
        page.insert_text((50, 60), f"{quarter} sales by region (units)", fontsize=22)
        page.insert_text((50, 90), "See the chart below.", fontsize=12)
        page.insert_image(pymupdf.Rect(50, 110, 792, 525), stream=_chart_png(values))
    doc.save(work / "sales.pdf")
    doc.close()


def _contract(path: Path, amount2: str, days: str, extra: bool) -> None:
    import docx

    d = docx.Document()
    d.add_heading("Service agreement", 1)
    d.add_paragraph("1. The Supplier provides hosting services to the Client.")
    d.add_paragraph(f"2. Invoices are payable within {days} of receipt.")
    d.add_paragraph("3. Payment schedule:")
    table = d.add_table(rows=1, cols=3)
    for cell, text in zip(
        table.rows[0].cells, ("Milestone", "Amount (EUR)", "Due"), strict=False
    ):
        cell.text = text
    for row in (
        ("Setup", "12,000", "March"),
        ("Migration", amount2, "June"),
        ("Go-live", "20,000", "September"),
    ):
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=False):
            cell.text = text
    d.add_paragraph("4. Governing law: Quebec.")
    if extra:
        d.add_paragraph(
            "5. Either party may terminate this agreement with 90 days' notice."
        )
    d.save(path)


def build_contracts(work: Path) -> None:
    _contract(work / "contract_v1.docx", "40,000", "30 days", False)
    _contract(work / "contract_v2.docx", "48,000", "60 days", True)


def build_repo(work: Path) -> None:
    repo = work / "repo"
    (repo / "src" / "client").mkdir(parents=True)
    (repo / "build" / "lib" / "client").mkdir(parents=True)
    (repo / "src" / "client" / "__init__.py").write_text("from .http import fetch\n")
    (repo / "src" / "client" / "config.py").write_text(
        '"""Client settings."""\n\nDEFAULT_RETRIES = 7\nDEFAULT_TIMEOUT = 12.5\n'
    )
    (repo / "src" / "client" / "http.py").write_text(
        "from .config import DEFAULT_RETRIES, DEFAULT_TIMEOUT\n\n\n"
        "def fetch(url, retries=DEFAULT_RETRIES):\n"
        "    for attempt in range(retries + 1):\n"
        "        ...  # one try, then `retries` more\n"
    )
    # A stale build output, ignored by git: its value is out of date.
    (repo / "build" / "lib" / "client" / "config.py").write_text(
        '"""Client settings."""\n\nDEFAULT_RETRIES = 3\nDEFAULT_TIMEOUT = 30\n'
    )
    (repo / ".gitignore").write_text("build/\n__pycache__/\n")
    (repo / "README.md").write_text(
        "# client\nA small HTTP client. Retries are configurable.\n"
    )
    env = {
        "GIT_AUTHOR_NAME": "dev",
        "GIT_AUTHOR_EMAIL": "d@x",
        "GIT_COMMITTER_NAME": "dev",
        "GIT_COMMITTER_EMAIL": "d@x",
        "PATH": "/run/current-system/sw/bin:/usr/bin:/bin",
    }
    for cmd in (
        ["git", "init", "-q", "-b", "main"],
        ["git", "add", "-A"],
        ["git", "commit", "-q", "-m", "init"],
    ):
        subprocess.run(cmd, cwd=repo, check=True, env={**env, **_path_env()})


def _path_env() -> dict:
    import os

    return {"PATH": os.environ.get("PATH", "")}


def mcp_report_pages() -> list[str]:
    rng = random.Random(17)
    pages = []
    for n in range(1, 41):
        sentences = [f"Project report, page {n}."]
        if n == 27:
            sentences.append("The approved budget for phase 3 is 2.4 million euros.")
        while sum(len(s) + 1 for s in sentences) < 1400:
            a, b = rng.sample(WORDS, 2)
            sentences.append(
                f"The {a} team reviewed the {b} plan ({n}-{len(sentences)})."
            )
        pages.append(" ".join(sentences))
    return pages


def build_mcp_report(work: Path) -> None:
    _text_pdf(work / "report.pdf", mcp_report_pages(), fontsize=9)


def build_mcp_folder(work: Path) -> None:
    p = work / "project"
    for d in ("src", "node_modules/left-pad", "secrets", "logs"):
        (p / d).mkdir(parents=True)
    (p / "src" / "main.py").write_text("print('hello')\n")
    (p / "src" / "util.py").write_text("def add(a, b):\n    return a + b\n")
    (p / "README.md").write_text("# Demo\n")
    (p / ".env").write_text("TOKEN=abc\n")
    (p / "node_modules" / "left-pad" / "index.js").write_text("module.exports = 1\n")
    (p / ".gitignore").write_text("*.log\nsecrets/\n")
    (p / "logs" / "app.log").write_text("started\n")
    (p / "secrets" / "key.txt").write_text("k\n")
    (p / ".hidden_notes.txt").write_text("h\n")
    (p / "poetry.lock").write_text("# lock\n")


def build_pages_pdf(work: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    for n in range(1, 5):
        page = doc.new_page()
        page.insert_text((72, 100), f"Hello from page {n} (PAGE-{n})", fontsize=20)
        page.draw_rect(pymupdf.Rect(72, 200, 300, 400), color=(0, 0, 1), fill=(0, 0, 1))
    doc.save(work / "report.pdf")
    doc.close()


def build_nothing(work: Path) -> None:
    pass


SITE_PAGE = """<!doctype html>
<html><head><title>Rivest Hardware — Spring price list</title></head>
<body>
<header>
  <nav><a href="/">Home</a> · <a href="/shop/">Shop</a> · <a href="/blog/">Blog</a></nav>
  <div class="banner">Flash deal: Widget Pro only $9.99 today with code FLASH!</div>
</header>
<main><article>
  <h1>Spring price list</h1>
  <p>Our list prices for the spring season, valid from March 1 until further notice.
     Prices are per unit, before tax.</p>
  <table>
    <thead><tr><th>Product</th><th>List price</th><th>Stock</th></tr></thead>
    <tbody>
      <tr><td>Widget Pro</td><td>$14.50</td><td>In stock</td></tr>
      <tr><td>Gadget Max</td><td>$22.00</td><td>Low</td></tr>
      <tr><td>Sprocket</td><td>$3.75</td><td>In stock</td></tr>
    </tbody>
  </table>
  <p>Bulk orders over 500 units get 8% off the list price.</p>
</article></main>
<footer>
  <p>Last year's list: Widget Pro $12.00, Gadget Max $19.00.</p>
  <p>Sitemap · Careers · Press · © Rivest Hardware</p>
</footer>
</body></html>
"""


def build_site(root: Path) -> None:
    (root / "shop").mkdir(parents=True, exist_ok=True)
    (root / "shop" / "index.html").write_text(SITE_PAGE)


# =============================================================================
# Graders (Python source run in the workspace; prelude imports grading.py)
# =============================================================================

PRELUDE = f"import sys; sys.path.insert(0, {str(HERE)!r}); from grading import *\n"

G_SHEET = r"""
m = load("context.py")
text = m.context()
emit({
    "is_text": isinstance(text, str),
    "first_row": "S-00001" in text,
    "middle_row": "S-02500" in text,
    "last_row": "S-05000" in text,
    "no_other_sheets": "SUMMARY-ONLY" not in text and "COSTS-ONLY" not in text,
    "info_uses_attachments": uses("context.py", "attachments"),
})
"""

G_CACHE = r"""
import lm15
m = load("ask.py")
r1, r2 = m.request("How many vacation days QX1?"), m.request("Can I work remotely QX2?")
c = {"is_lm15_request": isinstance(r1, lm15.Request) and isinstance(r2, lm15.Request)}
body = provider_bodies(r1)["anthropic"]
body2 = provider_bodies(r2)["anthropic"]
def blocks(b):
    out = []
    system = b.get("system")
    if isinstance(system, str):
        out.append({"type": "text", "text": system})
    elif isinstance(system, list):
        out += system
    for msg in b["messages"]:
        content = msg["content"]
        out += [{"type": "text", "text": content}] if isinstance(content, str) else content
    return out
b1, b2 = blocks(body), blocks(body2)
dump = json.dumps(b1) + pdf_text_in(b1)  # text, or the PDF itself
c["whole_handbook_sent"] = all(f"DOC-P{n}" in dump for n in (1, 2, 3))
c["question_sent"] = "QX1" in dump
doc_end = max((i for i, b in enumerate(b1)
               if "DOC-P3" in json.dumps(b) + pdf_text_in(b)), default=-1)
q_at = min((i for i, b in enumerate(b1) if "QX1" in json.dumps(b)), default=-1)
marks = [i for i, b in enumerate(b1) if isinstance(b, dict) and b.get("cache_control")]
c["cache_mark_after_handbook_before_question"] = any(doc_end <= i < q_at for i in marks)
cut = max([i for i in marks if doc_end <= i < q_at], default=-1)
c["same_prefix_for_every_question"] = cut >= 0 and json.dumps(b1[:cut + 1]) == json.dumps(b2[:cut + 1])
emit(c)
"""

G_PROVIDERS = r"""
import lm15
m = load("ask.py")
c = {"is_lm15_request": True, "question_everywhere": True, "pages_3_to_5_everywhere": True}
for provider, model in (("anthropic", "claude-sonnet-4-5"), ("openai", "gpt-4.1-mini"),
                        ("gemini", "gemini-2.5-flash")):
    req = m.request("What is the tallest bar on page 4?", model)
    c["is_lm15_request"] &= isinstance(req, lm15.Request)
    body = provider_bodies(req)[provider]
    c["question_everywhere"] &= "tallest bar on page 4" in json.dumps(body)
    pages = sorted(chart_page(img) for img in images_in(body))
    c[f"info_{provider}_pages"] = pages
    c["pages_3_to_5_everywhere"] &= pages == [3, 4, 5]
emit(c)
"""

G_LABELS = r"""
import io, re, lm15
m = load("build.py")
msg = m.message()
c = {"is_lm15_message": isinstance(msg, lm15.Message)}
imgs = []
for p in getattr(msg, "parts", ()):
    if getattr(p, "type", None) == "image":
        from PIL import Image
        imgs.append(Image.open(io.BytesIO(p.bytes)))
c["three_pictures"] = len(imgs) == 3
found = []
for img in imgs:
    text = ocr_text(img)
    found.append(sorted(set(int(n) for n in re.findall(r"\b([1-8])\b", text))))
c["info_numbers_read"] = found
# Its own number, and not the other two (a box drawn around a label can
# read as a stray "1": that is not a wrong label).
c["each_picture_shows_its_slide"] = len(found) == 3 and all(
    n in got and not (set(got) & ({2, 5, 7} - {n})) for n, got in zip((2, 5, 7), found)
)
emit(c)
"""

G_RAG = r"""
import re, subprocess, sys
sys.path.insert(0, %(here)r)
from tasks import manual_pages
p = subprocess.run([sys.executable, "index.py"], capture_output=True, text=True, timeout=600)
norm = lambda s: re.sub(r"\s+", " ", s).strip()
pages = [norm(t) for t in manual_pages()]
c = {"runs": p.returncode == 0, "valid_records": False, "small_chunks": False,
     "each_chunk_from_its_page": False, "covers_every_page": False}
try:
    rows = [json.loads(l) for l in open("chunks.jsonl") if l.strip()]
    c["valid_records"] = bool(rows) and all(
        isinstance(r.get("text"), str) and isinstance(r.get("page"), int) for r in rows)
    c["small_chunks"] = c["valid_records"] and all(len(r["text"]) <= 1000 for r in rows)
    def body(t):
        return norm("\n".join(l for l in t.splitlines() if not l.startswith("## ")))
    ok, covered = True, set()
    for r in rows:
        n = r["page"]
        if not 1 <= n <= 60 or body(r["text"]) not in pages[n - 1]:
            ok = False
        elif f"[P-{n:03d}]" in r["text"]:
            covered.add(n)
    # every sentence of every page somewhere: count characters per page
    per_page = {}
    for r in rows:
        if 1 <= r["page"] <= 60:
            per_page[r["page"]] = per_page.get(r["page"], 0) + len(body(r["text"]))
    full = all(per_page.get(n, 0) >= len(pages[n - 1]) * 0.98 for n in range(1, 61))
    c["each_chunk_from_its_page"] = c["valid_records"] and ok
    c["covers_every_page"] = len(covered) == 60 and full
    c["info_chunks"] = len(rows)
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
""" % {"here": str(HERE)}

G_SECRETS = r"""
import lm15
m = load("review.py")
msg = m.message()
dump = json.dumps([getattr(p, "text", "") for p in getattr(msg, "parts", ())])
emit({
    "is_lm15_message": isinstance(msg, lm15.Message),
    "code_and_docs_included": all(k in dump for k in
        ("CODE-SERVER", "CODE-BILLING", "README-MARKER", "SETTINGS-MARKER")),
    "no_secret_sent": "SECRET" not in dump and "PRIVATE KEY" not in dump,
})
"""

G_INVOICES = r"""
from pathlib import Path
EXPECTED = %(expected)r
c = {"uses_lm15": uses("totals.py", "lm15"), "all_totals_right": False}
try:
    got = json.loads(Path("totals.json").read_text())
    got = {Path(k).name: number(v) for k, v in got.items()}
    wrong = [k for k, v in EXPECTED.items() if got.get(k) is None or abs(got[k] - v) > 0.01]
    c["all_totals_right"] = not wrong
    c["info_wrong"] = wrong
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
""" % {"expected": {k: v for k, (v, _) in INVOICES.items()}}

G_CHART_VALUES = r"""
from pathlib import Path
QUARTERS = %(q)r
c = {"uses_lm15": uses("ask.py", "lm15"), "all_quarters_right": False}
try:
    got = json.loads(Path("answers.json").read_text())
    wrong = []
    for q, values in QUARTERS.items():
        region = max(values, key=values.get)
        a = got.get(q) or {}
        if str(a.get("region", "")).strip().lower() != region.lower() or number(a.get("value")) != values[region]:
            wrong.append(q)
    c["all_quarters_right"] = not wrong
    c["info_wrong"] = wrong
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
""" % {"q": QUARTERS}

G_CONTRACTS = r"""
import re
from pathlib import Path
c = {"uses_lm15": uses("compare.py", "lm15")}
try:
    text = Path("changes.md").read_text().lower()
    c["amount_change"] = bool(re.search(r"48[ ,.]?000", text))
    c["payment_days_change"] = "60" in text
    c["termination_clause"] = "terminat" in text and "90" in text
except Exception as exc:
    c["error"] = repr(exc)[:300]
    c["amount_change"] = False
emit(c)
"""

G_WEB = r"""
from pathlib import Path
c = {"uses_lm15": uses("ask.py", "lm15"), "price_right": False}
try:
    c["price_right"] = number(json.loads(Path("answer.json").read_text()).get("price")) == 14.5
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
"""

G_REPO = r"""
from pathlib import Path
# The files may be next to repo/ or inside it: the prompt does not say.
c = {"uses_lm15": uses(first_existing("ask.py", "repo/ask.py"), "lm15"), "retries_right": False}
try:
    answer = first_existing("answer.json", "repo/answer.json")
    c["retries_right"] = number(json.loads(Path(answer).read_text()).get("retries")) == 7
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
"""

G_MCP_PDF = r"""
from pathlib import Path
c = {"budget_right": False, "page_right": False}
try:
    a = json.loads(Path("answer.json").read_text())
    c["budget_right"] = "2.4" in str(a.get("budget", "")) and "million" in str(a.get("budget", "")).lower()
    c["page_right"] = number(a.get("page")) == 27
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
"""

G_MCP_FOLDER = r"""
from pathlib import Path
SKIPPED = {".env", "node_modules/left-pad/index.js", "logs/app.log", "secrets/key.txt",
           ".hidden_notes.txt", ".gitignore", "poetry.lock"}
READ = {"src/main.py", "src/util.py", "README.md"}
c = {"lists_every_skipped_file": False, "lists_no_read_file": False}
try:
    listed = [str(x).strip().removeprefix("./").removeprefix("project/")
              for x in json.loads(Path("skipped.json").read_text())]
    def covers(entry, path):
        entry = entry.rstrip("/")
        return path == entry or path.startswith(entry + "/")
    c["lists_every_skipped_file"] = all(any(covers(e, p) for e in listed) for p in SKIPPED)
    c["lists_no_read_file"] = not any(covers(e, p) for e in listed for p in READ)
    c["info_listed"] = listed
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
"""

G_CLI_JSON = r"""
import shutil, subprocess
from pathlib import Path
from attachments import Artifacts
c = {"no_python_written": not list(Path(".").glob("*.py")), "script_works": False,
     "loads_with_from_wire": False, "pages_2_3_only": False, "page_pictures": False}
work = Path.cwd()
copy = work.parent / "grading-cli"
shutil.rmtree(copy, ignore_errors=True)
shutil.copytree(work, copy, ignore=shutil.ignore_patterns(".venv", "pages.json"), symlinks=True)
(copy / ".venv").symlink_to(work / ".venv")
p = subprocess.run(["bash", "extract.sh"], cwd=copy, capture_output=True, text=True, timeout=300)
c["script_works"] = p.returncode == 0 and (copy / "pages.json").exists()
try:
    a = Artifacts.from_wire(json.loads((copy / "pages.json").read_text()))
    c["loads_with_from_wire"] = True
    t = a.text
    c["pages_2_3_only"] = "PAGE-2" in t and "PAGE-3" in t and "PAGE-1" not in t and "PAGE-4" not in t
    c["page_pictures"] = sorted(i.get("page") for i in a.images) == [2, 3]
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
"""

G_CLI_WEB = r"""
import re, shutil, subprocess
from pathlib import Path
c = {"no_python_written": not list(Path(".").glob("*.py")), "script_works": False,
     "article_text": False, "table_as_markdown": False, "no_menus_or_footer": False}
work = Path.cwd()
copy = work.parent / "grading-cli"
shutil.rmtree(copy, ignore_errors=True)
shutil.copytree(work, copy, ignore=shutil.ignore_patterns(".venv", "article.md"), symlinks=True)
(copy / ".venv").symlink_to(work / ".venv")
p = subprocess.run(["bash", "make_prompt.sh"], cwd=copy, capture_output=True, text=True, timeout=300)
c["script_works"] = p.returncode == 0 and (copy / "article.md").exists()
try:
    t = (copy / "article.md").read_text()
    c["article_text"] = "Spring price list" in t and "Bulk orders over 500" in t
    c["table_as_markdown"] = bool(re.search(r"^\|.*Widget Pro.*\|.*14\.50.*\|", t, re.M))
    c["no_menus_or_footer"] = not any(k in t for k in ("Flash deal", "Careers", "Last year's list", "<"))
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
"""


# =============================================================================
# The tasks
# =============================================================================


def _d(text: str) -> str:
    return " ".join(textwrap.dedent(text).split())


TASKS: dict[str, Task] = {}

# --- 1-5: from evals/skill -------------------------------------------------
for _name in ("photos", "charts", "folder", "install", "scans"):
    TASKS[_name] = Task(
        prompt=_d(skill_eval.PROMPTS[_name]),
        build=lambda work, n=_name: skill_eval.build_files(n, work),
        grader=None,
        expert=skill_eval.SOLUTIONS[_name],
        naive=skill_eval.NAIVE[_name],
        venv=_name != "install",
        naive_may_pass=(
            "the library fixed this trap in 1.0.0b3: scanned pages are JPEG by "
            "default, so the naive request is 22 MB, under the 32 MB limit"
            if _name == "scans"
            else ""
        ),
    )

# --- 6 ---------------------------------------------------------------------
TASKS["sheet"] = Task(
    prompt=_d(f"""
        data.xlsx has a sheet named Sales with 5,000 orders. A model will be asked
        questions that need every order. Write context.py with a function
        `context() -> str` that returns the text of the Sales sheet — every row, and
        nothing from the workbook's other sheets — to put in the model's prompt.
        {LIBS}"""),
    build=build_sheet,
    grader=G_SHEET,
    expert={
        "context.py": """
        from attachments import att

        def context() -> str:
            return att("data.xlsx[sheet: Sales, rows: 10000]").raise_for_errors().to_text(sources=False)
    """
    },
    naive={
        "context.py": """
        from attachments import att

        def context() -> str:
            return att("data.xlsx[sheet: Sales]").text
    """
    },
)

# --- 7-11: attachments + lm15, no model call --------------------------------
TASKS["cache"] = Task(
    prompt=_d(f"""
        We will ask Claude (model "claude-sonnet-4-5") many questions about
        handbook.pdf, one request per question. Write ask.py with a function
        `request(question: str) -> lm15.Request` that sends the whole handbook and
        the question, built so the provider can cache the handbook once and reuse it
        for every question (only the question changes from one request to the
        next). {LIBS} Do not call any API."""),
    build=build_handbook,
    grader=G_CACHE,
    expert={
        "ask.py": """
        from attachments import att
        from lm15 import CacheConfig, Config, Message, Request

        HANDBOOK = att("handbook.pdf").raise_for_errors().lm15()

        def request(question: str) -> Request:
            return Request(
                model="claude-sonnet-4-5",
                messages=(HANDBOOK, Message.user(question)),
                config=Config(cache=CacheConfig(prefix_until_index=0)),
            )
    """
    },
    naive={
        "ask.py": """
        from attachments import att
        from lm15 import CacheConfig, Config, Request

        def request(question: str) -> Request:
            return Request(model="claude-sonnet-4-5",
                           messages=(att("handbook.pdf").lm15(question),),
                           config=Config(cache=CacheConfig()))
    """
    },
)

TASKS["providers"] = Task(
    prompt=_d(f"""
        report.pdf is a 6-page slide deck; the numbers that matter are only in the
        bar charts on pages 3 to 5. Write ask.py with a function
        `request(question: str, model: str) -> lm15.Request` that lets a vision
        model answer questions about those charts. It must work unchanged for
        model "claude-sonnet-4-5", "gpt-4.1-mini" and "gemini-2.5-flash": we send it
        through lm15 to Anthropic, OpenAI and Gemini. {LIBS} Do not call any API."""),
    build=build_charts,
    grader=G_PROVIDERS,
    expert={
        "ask.py": """
        from attachments import att
        from lm15 import Request

        def request(question: str, model: str) -> Request:
            a = att("report.pdf[pages: 3-5, images: true, max_dim: 1568]").raise_for_errors()
            return Request(model=model, messages=(a.lm15(question),))
    """
    },
    naive={
        "ask.py": """
        from attachments import att
        from lm15 import Request

        def request(question: str, model: str) -> Request:
            return Request(model=model, messages=(att("report.pdf[pages: 3-5]").lm15(question),))
    """
    },
)

TASKS["labels"] = Task(
    prompt=_d(f"""
        deck.pdf has 8 slides with charts. A vision model will get pictures of
        slides 2, 5 and 7 and must say which slide each chart is on, from the
        pictures alone: our pipeline may drop the text around them. Write build.py
        with a function `message() -> lm15.Message` holding those three slides as
        pictures, in that order, each picture showing its own slide number.
        {LIBS}"""),
    build=build_deck,
    grader=G_LABELS,
    expert={
        "build.py": """
        from attachments import att

        def message():
            return att("deck.pdf[pages: 2,5,7, images: true, watermark: page]").raise_for_errors().lm15()
    """
    },
    naive={
        "build.py": """
        from attachments import att

        def message():
            return att("deck.pdf[pages: 2,5,7, images: true]").lm15()
    """
    },
)

TASKS["rag"] = Task(
    prompt=_d("""
        manual.pdf is a 60-page product manual. For our search index, write
        index.py: running `.venv/bin/python index.py` writes chunks.jsonl, one JSON
        object per line, {"text": ..., "page": ...}: chunks of at most 1000
        characters, each taken from exactly one page, `page` being that page's
        number (1-based), together covering the whole manual. Use the Python
        library attachments, already installed in .venv."""),
    build=build_manual,
    grader=G_RAG,
    expert={
        "index.py": """
        import json
        from attachments import Artifacts, att, make_artifact

        a = att("manual.pdf").raise_for_errors()[0]
        with open("chunks.jsonl", "w") as f:
            for seg in a["meta"]["segments"]:
                page = Artifacts([make_artifact(text=a["text"][seg["start"]:seg["end"]])])
                for text in page.chunk(max_chars=1000, overlap=0, sources=False):
                    f.write(json.dumps({"text": text, "page": seg["page"]}) + "\\n")
    """
    },
    naive={
        "index.py": """
        import json
        from attachments import att

        with open("chunks.jsonl", "w") as f:
            for i, text in enumerate(att("manual.pdf").chunk(max_chars=1000)):
                f.write(json.dumps({"text": text, "page": i + 1}) + "\\n")
    """
    },
)

TASKS["secrets"] = Task(
    prompt=_d(f"""
        proj/ is a small codebase. Write review.py with a function
        `message() -> lm15.Message` holding the project (its code, docs and
        configuration) for an outside model provider that will review it.
        {LIBS} Do not call any API."""),
    build=build_project,
    grader=G_SECRETS,
    expert={
        "review.py": """
        from attachments import att

        def message():
            # The folder rules skip .env by name; a key in a file without a
            # telling name is not recognised: look at the overview, leave it out.
            return att("proj/[ignore: deploy_key]").raise_for_errors().lm15("Review this project.")
    """
    },
    naive={
        "review.py": """
        from pathlib import Path
        from lm15 import Message

        def message():
            files = [p for p in sorted(Path("proj").rglob("*")) if p.is_file()]
            return Message.user("\\n\\n".join(f"## {p}\\n{p.read_text()}" for p in files))
    """
    },
)

# --- 12-16: real model calls through lm15 ----------------------------------
TASKS["invoices"] = Task(
    prompt=_d(f"""
        invoices/ holds 6 supplier invoices as PDFs; some are scans. Write
        totals.py: running it asks a model for each invoice's total amount due and
        writes totals.json, mapping each file name to its total as a number. Run it
        so totals.json exists. {CALL} {LIBS}"""),
    build=build_invoices,
    grader=G_INVOICES,
    calls_model=True,
    expert={
        "totals.py": f"""
        import json, re
        from pathlib import Path
        from attachments import att
        from lm15 import LMRouter, Request

        router = LMRouter()
        totals = {{}}
        for pdf in sorted(Path("invoices").glob("*.pdf")):
            a = att(f"{{pdf}}[images: true]").raise_for_errors()
            reply = router.complete(Request(model="{TASK_MODEL}", messages=(
                a.lm15("What is the total amount due? Reply with the number only."),)))
            totals[pdf.name] = float(re.sub(r"[^0-9.]", "", reply.text.replace(",", "")))
        Path("totals.json").write_text(json.dumps(totals, indent=1))
    """,
        "__run__": ["totals.py"],
    },
    naive={
        "totals.py": f"""
        import json, re
        from pathlib import Path
        from attachments import att
        from lm15 import LMRouter, Request

        router = LMRouter()
        totals = {{}}
        for pdf in sorted(Path("invoices").glob("*.pdf")):
            text = att(f"{{pdf}}[images: false, ocr: false]").to_text()
            reply = router.complete(Request(model="{TASK_MODEL}", messages=(
                __import__("lm15").Message.user(text + "\\nTotal amount due? Number only; 0 if unknown."),)))
            totals[pdf.name] = float(re.sub(r"[^0-9.]", "", reply.text.replace(",", "")) or 0)
        Path("totals.json").write_text(json.dumps(totals))
    """,
        "__run__": ["totals.py"],
    },
)

TASKS["chart_values"] = Task(
    prompt=_d(f"""
        sales.pdf has one page per quarter with a bar chart of units sold by
        region. Write ask.py: running it asks a model, for each quarter, which
        region sold the most and how many units, and writes answers.json as
        {{"Q1": {{"region": "...", "value": <units>}}, ..., "Q4": {{...}}}}. Run it.
        {CALL} {LIBS}"""),
    build=build_sales,
    grader=G_CHART_VALUES,
    calls_model=True,
    expert={
        "ask.py": f"""
        import json
        from pathlib import Path
        from attachments import att
        from lm15 import Config, LMRouter, Request

        a = att("sales.pdf[images: true]").raise_for_errors()
        prompt = ('For each page (Q1 to Q4), which region sold the most and how many units? '
                  'Reply with JSON only: {{"Q1": {{"region": "...", "value": 0}}, ...}}')
        reply = LMRouter().complete(Request(model="{TASK_MODEL}", messages=(a.lm15(prompt),)))
        text = reply.text.strip().removeprefix("```json").removesuffix("```")
        Path("answers.json").write_text(json.dumps(json.loads(text), indent=1))
    """,
        "__run__": ["ask.py"],
    },
    naive={
        "ask.py": f"""
        import json
        from pathlib import Path
        from attachments import att
        from lm15 import LMRouter, Request

        a = att("sales.pdf")
        prompt = ('For each quarter, which region sold the most and how many units? JSON only: '
                  '{{"Q1": {{"region": "...", "value": 0}}, ...}}; use "unknown" and 0 if you cannot tell.')
        reply = LMRouter().complete(Request(model="{TASK_MODEL}", messages=(a.lm15(prompt),)))
        text = reply.text.strip().removeprefix("```json").removesuffix("```")
        Path("answers.json").write_text(text)
    """,
        "__run__": ["ask.py"],
    },
)

TASKS["contracts"] = Task(
    prompt=_d(f"""
        contract_v1.docx and contract_v2.docx are two versions of a service
        agreement. Write compare.py: running it asks a model what changed from
        version 1 to version 2 (every change, nothing invented) and writes the
        answer to changes.md. Run it. {CALL} {LIBS}"""),
    build=build_contracts,
    grader=G_CONTRACTS,
    calls_model=True,
    expert={
        "compare.py": f"""
        from pathlib import Path
        from attachments import att
        from lm15 import LMRouter, Request

        a = att(["contract_v1.docx", "contract_v2.docx"]).raise_for_errors()
        reply = LMRouter().complete(Request(model="{TASK_MODEL}", messages=(a.lm15(
            "List every change from contract_v1.docx to contract_v2.docx, tables included. "
            "Only real changes."),)))
        Path("changes.md").write_text(reply.text)
    """,
        "__run__": ["compare.py"],
    },
    naive={
        "compare.py": f"""
        from pathlib import Path
        import docx
        from lm15 import LMRouter, Message, Request

        def text(path):
            return "\\n".join(p.text for p in docx.Document(path).paragraphs)
        prompt = f"V1:\\n{{text('contract_v1.docx')}}\\n\\nV2:\\n{{text('contract_v2.docx')}}\\n\\nWhat changed?"
        reply = LMRouter().complete(Request(model="{TASK_MODEL}", messages=(Message.user(prompt),)))
        Path("changes.md").write_text(reply.text)
    """,
        "__run__": ["compare.py"],
    },
)

TASKS["web"] = Task(
    prompt=_d(f"""
        Our shop's price list is the web page {{site}}/shop/. Write ask.py: running
        it gives that page to a model and asks the current list price of Widget
        Pro, then writes answer.json as {{"price": <number>}}. Run it. {CALL}
        {LIBS}"""),
    build=build_nothing,
    grader=G_WEB,
    calls_model=True,
    site=True,
    expert={
        "ask.py": f"""
        import json, re, sys
        from pathlib import Path
        from attachments import att
        from lm15 import LMRouter, Request

        a = att(sys.argv[1] + "/shop/").raise_for_errors()
        reply = LMRouter().complete(Request(model="{TASK_MODEL}", messages=(a.lm15(
            "What is the current list price of Widget Pro? Number only."),)))
        Path("answer.json").write_text(json.dumps({{"price": float(re.sub(r"[^0-9.]", "", reply.text))}}))
    """,
        "__run__": ["ask.py", "{site}"],
    },
    naive={
        "ask.py": """
        import json
        from pathlib import Path
        Path("answer.json").write_text(json.dumps({"price": 9.99}))
    """
    },
)

TASKS["repo"] = Task(
    prompt=_d(f"""
        repo/ is a Python project. Write ask.py: running it gives the project to a
        model and asks how many times the client retries a failed request by
        default, then writes answer.json as {{"retries": <integer>}}. Run it.
        {CALL} {LIBS}"""),
    build=build_repo,
    grader=G_REPO,
    calls_model=True,
    expert={
        "ask.py": f"""
        import json, re
        from pathlib import Path
        from attachments import att
        from lm15 import LMRouter, Request

        a = att("repo/").raise_for_errors()
        reply = LMRouter().complete(Request(model="{TASK_MODEL}", messages=(a.lm15(
            "How many times does the client retry a failed request by default? Integer only."),)))
        Path("answer.json").write_text(json.dumps({{"retries": int(re.sub(r"[^0-9]", "", reply.text))}}))
    """,
        "__run__": ["ask.py"],
    },
    naive={
        "ask.py": f"""
        import json, re
        from pathlib import Path
        from attachments import att
        from lm15 import LMRouter, Request

        a = att("repo/README.md")
        reply = LMRouter().complete(Request(model="{TASK_MODEL}", messages=(a.lm15(
            "How many times does the client retry by default? Integer only; 0 if not stated."),)))
        Path("answer.json").write_text(json.dumps({{"retries": int(re.sub(r"[^0-9]", "", reply.text) or 0)}}))
    """,
        "__run__": ["ask.py"],
    },
)

# --- 17-18: MCP only -------------------------------------------------------
TASKS["mcp_pdf"] = Task(
    prompt=_d(f"""
        {MCP} report.pdf is a 40-page project report. Find the approved budget for
        phase 3 and the page it is on, and write answer.json as
        {{"budget": "...", "page": <page number>}}."""),
    build=build_mcp_report,
    grader=G_MCP_PDF,
    mcp=True,
    expert={"answer.json": '{"budget": "2.4 million euros", "page": 27}'},
    naive={"answer.json": '{"budget": "2.4 million euros", "page": 26}'},
)

TASKS["mcp_folder"] = Task(
    prompt=_d(f"""
        {MCP} Find out which files of project/ the attachments library leaves out
        when it reads that folder, and write skipped.json: a JSON list of their
        paths relative to project/ (a folder path ending in "/" stands for
        everything in it)."""),
    build=build_mcp_folder,
    grader=G_MCP_FOLDER,
    mcp=True,
    expert={
        "skipped.json": json.dumps(
            [
                ".env",
                "node_modules/",
                "logs/app.log",
                "secrets/",
                ".hidden_notes.txt",
                ".gitignore",
                "poetry.lock",
            ]
        )
    },
    naive={"skipped.json": json.dumps([".env", "node_modules/"])},
)

# --- 19-20: the command line ------------------------------------------------
TASKS["cli_json"] = Task(
    prompt=_d("""
        Without writing any Python, using the `att` command (in .venv/bin), write
        extract.sh: it saves pages 2 and 3 of report.pdf, with a picture of each of
        those pages, to pages.json, in the form that the attachments library's
        `Artifacts.from_wire` loads. Run it."""),
    build=build_pages_pdf,
    grader=G_CLI_JSON,
    expert={
        "extract.sh": """
        #!/usr/bin/env bash
        set -euo pipefail
        .venv/bin/att "report.pdf[pages: 2-3, images: true]" --json > pages.json
    """,
        "__run__sh__": ["extract.sh"],
    },
    naive={
        "extract.sh": """
        #!/usr/bin/env bash
        .venv/bin/att report.pdf --json > pages.json
    """,
        "__run__sh__": ["extract.sh"],
    },
)

TASKS["cli_web"] = Task(
    prompt=_d("""
        Using only the `att` command (in .venv/bin), no Python, write
        make_prompt.sh: it saves the article at {site}/shop/ as Markdown to
        article.md — the article's text and its table, without the site's menus,
        banners and footer. Run it."""),
    build=build_nothing,
    grader=G_CLI_WEB,
    site=True,
    expert={
        "make_prompt.sh": """
        #!/usr/bin/env bash
        set -euo pipefail
        .venv/bin/att "{site}/shop/" > article.md
    """,
        "__run__sh__": ["make_prompt.sh"],
    },
    naive={
        "make_prompt.sh": """
        #!/usr/bin/env bash
        curl -s "{site}/shop/" > article.md
    """,
        "__run__sh__": ["make_prompt.sh"],
    },
)

assert len(TASKS) == 20, len(TASKS)


# =============================================================================
# The same tasks without attachments (run.py --arm none)
# =============================================================================
#
# The workspace has lm15 and the libraries attachments itself uses (PyMuPDF,
# pypdf, python-docx, openpyxl, pandas, Pillow, RapidOCR, BeautifulSoup), so
# the only difference is the attachments library. Tasks that are about
# attachments itself (install, mcp_folder, cli_json) have no counterpart.

#: What the no-attachments workspace has installed (same versions uv picks
#: for attachments' own extras).
NONE_PACKAGES = (
    "lm15>=1.2",
    "pymupdf>=1.24",
    "pypdf>=6.17",
    "python-docx>=1.0",
    "openpyxl>=3.1,<4.0",
    "pandas>=2.0,<3.0",
    "xlrd>=2.0",
    "pillow>=10.1",
    "beautifulsoup4>=4.12",
    "lxml>=5.0",
    "rapidocr>=3.10,<4",
    "onnxruntime>=1.18",
)

LIBS_NONE = (
    "Use Python from .venv (run it as .venv/bin/python): it has lm15 and "
    "common libraries for reading PDFs, Office files, images, OCR and web "
    "pages (`uv pip list --python .venv/bin/python` lists them); add others "
    "with `uv pip install --python .venv/bin/python <package>` if you need to. "
    "Do not use the attachments library."
)

_SWAPS = (
    (LIBS, LIBS_NONE),
    (
        "Use the Python library attachments, already installed in .venv (run Python "
        "as .venv/bin/python).",
        LIBS_NONE,
    ),
    ("Use the Python library attachments, already installed in .venv.", LIBS_NONE),
    (MCP + " ", ""),
)

NONE_PROMPTS: dict[str, str] = {}
for _name, _task in TASKS.items():
    if _name in ("install", "mcp_folder", "cli_json", "folder", "cli_web"):
        continue
    _p = _task.prompt
    for _old, _new in _SWAPS:
        _p = _p.replace(_old, _new)
    if _name == "mcp_pdf":
        _p += " " + LIBS_NONE
    NONE_PROMPTS[_name] = _p

NONE_PROMPTS["folder"] = _d(f"""
    docs/ holds a project's documents (PDF, Word, Excel, text, a picture).
    Write prepare.py. Running `.venv/bin/python prepare.py` must read every
    file in docs/ and save the result to docs.json, pictures included, so a
    later program can get it back without reading the documents again: give
    prepare.py a function `load()` that returns it as a list with one dict per
    file: {{"name": <file name>, "text": <its text>, "images": [<the bytes of
    each picture>]}}. If any document cannot be read, prepare.py must exit
    with a non-zero status, print an error naming every unreadable file, and
    write nothing. {LIBS_NONE}""")

NONE_PROMPTS["cli_web"] = _d(f"""
    Write make_prompt.sh: it saves the article at {{site}}/shop/ as Markdown to
    article.md — the article's text and its table, without the site's menus,
    banners and footer. Run it. {LIBS_NONE}""")

for _name, _p in NONE_PROMPTS.items():
    assert "attachments" not in _p.replace("Do not use the attachments library.", ""), (
        _name
    )

G_FOLDER_NONE = r"""
import importlib.util, io, json, os, shutil, subprocess, sys
from pathlib import Path
from PIL import Image
py = sys.executable
work = Path.cwd()
c = {"stops_naming_all": False, "writes_nothing_on_failure": False,
     "saves_and_loads": False, "load_needs_no_documents": False}
def run(folder):
    p = subprocess.run([py, "prepare.py"], cwd=folder, capture_output=True, text=True, timeout=300)
    return p.returncode, p.stdout + p.stderr
grading = work.parent / "grading"
shutil.rmtree(grading, ignore_errors=True)
for name in ("broken", "clean"):
    shutil.copytree(work, grading / name, ignore=shutil.ignore_patterns(".venv", "docs.json", "__pycache__"))
code, out = run(grading / "broken")
c["stops_naming_all"] = code != 0 and "broken.pdf" in out and "damaged.docx" in out
c["writes_nothing_on_failure"] = code != 0 and not (grading / "broken" / "docs.json").exists()
for bad in ("broken.pdf", "damaged.docx"):
    (grading / "clean" / "docs" / bad).unlink()
code, out = run(grading / "clean")
EXPECTED = {"brief.pdf": "Project brief", "notes.txt": "Kickoff", "budget.xlsx": "servers",
            "memo.docx": "two more testers", "diagram.png": None}
os.chdir(grading / "clean")
try:
    json.loads(Path("docs.json").read_text())
    m = load("prepare.py")
    def intact(loaded):
        got = {Path(d["name"]).name: d for d in loaded}
        if set(got) != set(EXPECTED) or not got["diagram.png"]["images"]:
            return False
        for name, phrase in EXPECTED.items():
            if phrase and phrase not in got[name]["text"]:
                return False
        for d in loaded:
            for b in d["images"]:
                Image.open(io.BytesIO(b)).load()
        return True
    c["saves_and_loads"] = code == 0 and intact(m.load())
    shutil.move("docs", "docs-away")
    c["load_needs_no_documents"] = c["saves_and_loads"] and intact(m.load())
except Exception as exc:
    c["error"] = repr(exc)[:300]
emit(c)
"""

G_CLI_WEB_NONE = G_CLI_WEB.replace(
    'c = {"no_python_written": not list(Path(".").glob("*.py")), "script_works": False,',
    'c = {"script_works": False,',
)
assert G_CLI_WEB_NONE != G_CLI_WEB

#: Graders that differ without attachments (the others are outcome-only).
NONE_GRADERS = {"folder": G_FOLDER_NONE, "cli_web": G_CLI_WEB_NONE}

#: Expert solutions for the graders written for this arm (the others were
#: checked with the attachments arm and grade outcomes only).
NONE_EXPERTS = {
    "folder": {
        "prepare.py": """
        import base64, json, sys
        from pathlib import Path

        def read(path):
            if path.suffix == ".pdf":
                import pymupdf
                with pymupdf.open(path) as doc:
                    return "\\n".join(p.get_text() for p in doc), []
            if path.suffix == ".docx":
                import docx
                return "\\n".join(p.text for p in docx.Document(path).paragraphs), []
            if path.suffix == ".xlsx":
                import openpyxl
                book = openpyxl.load_workbook(path, read_only=True)
                return "\\n".join(",".join(str(v) for v in row) for s in book for row in s.iter_rows(values_only=True)), []
            if path.suffix == ".png":
                return "", [path.read_bytes()]
            return path.read_text(), []

        def load():
            data = json.loads(Path("docs.json").read_text())
            return [{**d, "images": [base64.b64decode(b) for b in d["images"]]} for d in data]

        if __name__ == "__main__":
            out, bad = [], []
            for path in sorted(Path("docs").iterdir()):
                try:
                    text, images = read(path)
                    out.append({"name": path.name, "text": text, "images": images})
                except Exception as e:
                    bad.append(f"{path.name}: {e}")
            if bad:
                sys.exit("cannot read:\\n" + "\\n".join(bad))
            Path("docs.json").write_text(json.dumps(
                [{**d, "images": [base64.b64encode(b).decode() for b in d["images"]]} for d in out]))
    """
    },
    "cli_web": {
        "make_prompt.sh": """
        #!/usr/bin/env bash
        set -euo pipefail
        .venv/bin/python - "{site}/shop/" > article.md <<'PY'
        import sys, urllib.request
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(urllib.request.urlopen(sys.argv[1]).read(), "lxml")
        art = soup.find("article")
        out = []
        for el in art.find_all(["h1", "p", "table"], recursive=True):
            if el.name == "h1":
                out.append("# " + el.get_text(strip=True))
            elif el.name == "p":
                out.append(" ".join(el.get_text().split()))
            else:
                rows = [[c.get_text(strip=True) for c in r.find_all(["th", "td"])] for r in el.find_all("tr")]
                out.append("| " + " | ".join(rows[0]) + " |")
                out.append("|" + " --- |" * len(rows[0]))
                out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        print("\\n\\n".join(out))
        PY
    """,
        "__run__sh__": ["make_prompt.sh"],
    },
}
