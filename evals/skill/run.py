#!/usr/bin/env python3
"""Does the attachments skill help a coding agent? Same tasks, with and without it.

Each run gives a fresh Pi agent one task in a fresh workspace and grades the
result by code (never by asking a model). The arms differ only in the skill:

    skill     pi --no-skills --skill <src/attachments/skill>
    none      pi --no-skills, and the skill folder removed from the
              installed package so the agent cannot find it there

Tasks (each comes from a trap seen in real use; see README.md):

    photos    Claude request bodies for naming the animal in 6 photos whose
              file names are the answers               -> names must not leak
    charts    OpenAI messages so a model can read charts on pages 3-5 of a
              PDF that also has text, pictures <= 1600 px -> images: true
    folder    read a folder into a reloadable JSON cache; stop and name every
              unreadable file, writing nothing          -> att() never raises
    install   an empty uv project: print pages 2-3 of a PDF -> PyPI has 0.25
    scans     one Claude request that reads a 40-page scanned PDF -> the PNG
              default makes a 126 MB request; Anthropic accepts 32 MB

Usage (from the repo root):

    uv run python evals/skill/run.py --check            # graders vs expert solutions
    uv run python evals/skill/run.py --reps 3 --parallel 6
    uv run python evals/skill/run.py --tasks photos --arms skill --reps 1
    uv run python evals/skill/run.py --report evals/skill/runs/<stamp>

Workspaces stay in ~/.cache/attachments-skill-eval/<stamp>/ for inspection;
results (grades, agent metrics, Pi sessions) go to evals/skill/runs/<stamp>/.
"""

from __future__ import annotations

import argparse
import concurrent.futures as futures
import json
import os
import re
import shutil
import subprocess
import sys
import textwrap
import time
from pathlib import Path

HERE = Path(__file__).resolve().parent
REPO = HERE.parents[1]
SKILL = REPO / "src" / "attachments" / "skill"
CACHE = Path.home() / ".cache" / "attachments-skill-eval"
TASKS = ("photos", "charts", "folder", "install", "scans")
ARMS = ("skill", "none")

RULES = (
    "You are working in the current folder for a user who will read your final "
    "message. Do not open, list or search anything under ~/Projects/attachments "
    "or ~/.cache/attachments-skill-eval outside the current folder. Finish the "
    "task completely, then reply with a short summary."
)
FORBIDDEN = re.compile(r"Projects/attachments|attachments/skill|SKILL\.md")

PROMPTS = {
    "photos": """\
        The folder photos/ holds 6 pictures; each file is named after the animal
        it shows (cat_1.jpg, dog_2.jpg, ...). We will ask a vision model to name
        the animal in each picture and score its answers against the file names.
        Write classify.py with a function `request(path: str) -> dict` that
        returns the complete JSON body of one call to Anthropic's Messages API
        (model "claude-sonnet-4-5", max_tokens 20) asking which animal the
        picture shows. Use the Python library attachments, already installed in
        .venv (run Python as .venv/bin/python). Do not call the API.""",
    "charts": """\
        report.pdf is a 6-page slide deck. The numbers that matter are only in
        the bar charts on pages 3 to 5. Write ask.py with a function
        `messages(question: str) -> list` that returns the `messages` argument
        of OpenAI's Chat Completions API (a vision model) so the model can
        answer questions about those charts. Requests must stay small: no
        picture may be larger than 1600 pixels on its longest side. Use the
        Python library attachments, already installed in .venv (run Python as
        .venv/bin/python). Do not call the API.""",
    "folder": """\
        docs/ holds a project's documents (PDF, Word, Excel, text, a picture).
        Write prepare.py. Running `.venv/bin/python prepare.py` must read every
        file in docs/ with the Python library attachments (already installed in
        .venv) and save the result to docs.json, pictures included, so a later
        program can get it back without reading the documents again: give
        prepare.py a function `load()` that returns it as the library's own
        result object. If any document cannot be read, prepare.py must exit
        with a non-zero status, print an error naming every unreadable file,
        and write nothing.""",
    "scans": """\
        scan.pdf is a 40-page scanned inspection report (pictures of pages, no
        text layer). Write ask.py with a function `request(question: str) ->
        dict` that returns the complete JSON body of ONE call to Anthropic's
        Messages API (model "claude-sonnet-4-5", max_tokens 1000) that lets
        Claude read every page to answer the question. It must be a request
        Anthropic accepts. Use the Python library attachments, already
        installed in .venv (run Python as .venv/bin/python). Do not call the
        API.""",
    "install": """\
        This is a new uv project with no dependencies yet. Using the Python
        library attachments, write show.py that prints the text of pages 2 and
        3 of report.pdf (and nothing from the other pages). Add what you need
        to the project with uv; `uv run python show.py` must work.""",
}


# =============================================================================
# Workspaces
# =============================================================================


def _pdf(
    path: Path,
    pages: list[tuple[str, tuple[float, float, float] | None]],
    size=(960, 540),
):
    import pymupdf

    doc = pymupdf.open()
    for title, color in pages:
        page = doc.new_page(width=size[0], height=size[1])
        page.insert_text((60, 70), title, fontsize=26)
        if color:  # a bar chart with no numbers as text: only a picture shows it
            for i, height in enumerate((120, 260, 180, 330)):
                x = 120 + i * 180
                rect = pymupdf.Rect(x, 480 - height, x + 110, 480)
                page.draw_rect(rect, color=color, fill=color)
    doc.save(path)
    doc.close()


def _scanned_pdf(path: Path, pages: int) -> None:
    """A realistic scan: 150 dpi greyscale photos of text pages, sensor noise,
    a slight tilt, embedded as JPEG (as scanners do). About 13 MB."""
    import io

    import numpy as np
    import pymupdf
    from PIL import Image, ImageDraw, ImageFilter

    rng = np.random.default_rng(0)
    doc = pymupdf.open()
    for n in range(1, pages + 1):
        img = Image.new("L", (1275, 1650), 236)
        draw = ImageDraw.Draw(img)
        draw.text((100, 80), f"INSPECTION REPORT - PAGE {n}", fill=20)
        for y in range(160, 1550, 28):
            draw.text(
                (100, y),
                f"Observed corrosion on bracket {y}, {y / 97:.1f} mm, action required.",
                fill=35,
            )
        noisy = np.array(img, dtype=np.float32) + rng.normal(0, 9, (1650, 1275))
        img = Image.fromarray(noisy.clip(0, 255).astype("uint8"))
        img = img.filter(ImageFilter.GaussianBlur(0.6)).rotate(0.4, fillcolor=236)
        buf = io.BytesIO()
        img.convert("RGB").save(buf, "JPEG", quality=80)
        page = doc.new_page(width=612, height=792)
        page.insert_image(page.rect, stream=buf.getvalue())
    doc.save(path)
    doc.close()


#: Bar colour per page of the charts deck; the grader names a picture's page by it.
CHART_COLORS = {
    1: (0.85, 0.1, 0.1),
    2: (0.1, 0.65, 0.1),
    3: (0.1, 0.2, 0.9),
    4: (1.0, 0.55, 0.0),
    5: (0.55, 0.1, 0.7),
    6: (0.4, 0.4, 0.4),
}
ANIMALS = ("cat", "dog", "horse", "owl", "fox", "cow")


def build_files(task: str, work: Path) -> None:
    from PIL import Image, ImageDraw

    work.mkdir(parents=True, exist_ok=True)
    if task == "photos":
        (work / "photos").mkdir()
        for i, animal in enumerate(ANIMALS, start=1):
            img = Image.new("RGB", (1600, 1200), (40 * i, 120, 200 - 25 * i))
            ImageDraw.Draw(img).ellipse((400, 300, 1200, 900), fill=(240, 220, 200))
            img.save(work / "photos" / f"{animal}_{i}.jpg", quality=90)
    elif task == "charts":
        _pdf(
            work / "report.pdf",
            [
                (f"Section {n}: regional sales (PAGE-{n})", CHART_COLORS[n])
                for n in range(1, 7)
            ],
        )
    elif task == "folder":
        docs = work / "docs"
        docs.mkdir()
        _pdf(
            docs / "brief.pdf",
            [("Project brief, page 1", None), ("Timeline, page 2", None)],
            (612, 792),
        )
        (docs / "notes.txt").write_text("Kickoff notes: ship the pilot in March.\n")
        import openpyxl

        book = openpyxl.Workbook()
        book.active.title = "Budget"
        book.active.append(["item", "cost"])
        book.active.append(["servers", 1200])
        book.save(docs / "budget.xlsx")
        import docx

        memo = docx.Document()
        memo.add_paragraph("Memo: the pilot needs two more testers.")
        memo.save(docs / "memo.docx")
        Image.new("RGB", (900, 600), (30, 90, 160)).save(docs / "diagram.png")
        (docs / "broken.pdf").write_bytes(b"%PDF-1.7\nthis file was cut off")
        (docs / "damaged.docx").write_bytes(b"PK\x03\x04 not really a zip")
    elif task == "scans":
        _scanned_pdf(work / "scan.pdf", 40)
    elif task == "install":
        _pdf(
            work / "report.pdf",
            [(f"Hello from page {n} (PAGE-{n})", None) for n in range(1, 5)],
            (612, 792),
        )
        (work / "pyproject.toml").write_text(
            '[project]\nname = "demo"\nversion = "0.1.0"\nrequires-python = ">=3.12"\n'
            "dependencies = []\n"
        )


def build_venv(work: Path, wheel: Path, arm: str) -> None:
    """A .venv with attachments 1.0 (not for `install`, which starts empty)."""
    env = {k: v for k, v in os.environ.items() if k != "VIRTUAL_ENV"}
    subprocess.run(
        ["uv", "venv", "-q", "--python", "3.12", str(work / ".venv")],
        check=True,
        env=env,
    )
    subprocess.run(
        [
            "uv",
            "pip",
            "install",
            "-q",
            "--python",
            str(work / ".venv" / "bin" / "python"),
            f"attachments[pdf,image,docx,xlsx] @ {wheel.as_uri()}",
        ],
        check=True,
        env=env,
    )
    if arm == "none":  # the skill ships in the package; this arm must not find it
        site = next((work / ".venv" / "lib").glob("python*/site-packages"))
        shutil.rmtree(site / "attachments" / "skill")


def workspace(task: str, arm: str, work: Path, wheel: Path) -> None:
    build_files(task, work)
    if task != "install":
        build_venv(work, wheel, arm)


# =============================================================================
# Graders (run with the workspace's own Python, so they see what the agent built)
# =============================================================================

GRADERS = {
    "photos": r"""
import base64, importlib.util, io, json, sys
from pathlib import Path
from PIL import Image
spec = importlib.util.spec_from_file_location("classify", "classify.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
c = {"valid_body": True, "image_sent": True, "no_name_leak": True}
for photo in sorted(Path("photos").iterdir()):
    body = m.request(str(photo))
    dumped = json.dumps(body)
    ok = (body.get("model") == "claude-sonnet-4-5" and body.get("max_tokens") == 20
          and isinstance(body.get("messages"), list) and body["messages"])
    c["valid_body"] &= bool(ok)
    images = [b for msg in body.get("messages", []) if isinstance(msg.get("content"), list)
              for b in msg["content"] if isinstance(b, dict) and b.get("type") == "image"]
    sent = False
    for b in images:
        try:
            Image.open(io.BytesIO(base64.b64decode(b["source"]["data"]))).load(); sent = True
        except Exception:
            pass
    c["image_sent"] &= sent
    c["no_name_leak"] &= photo.name not in dumped and photo.stem not in dumped
print(json.dumps(c))
""",
    "charts": r"""
import base64, importlib.util, io, json, re
from PIL import Image
COLORS = %(colors)s
spec = importlib.util.spec_from_file_location("ask", "ask.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
msgs = m.messages("What is the tallest bar on page 4?")
json.dumps(msgs)
c = {"valid_messages": True, "question_sent": False, "pages_3_to_5": False, "small": True}
pages, order = [], []
for msg in msgs:
    c["valid_messages"] &= msg.get("role") in ("system", "user", "assistant")
    content = msg.get("content")
    parts = [{"type": "text", "text": content}] if isinstance(content, str) else content
    for p in parts:
        if p.get("type") == "text":
            c["question_sent"] |= "tallest bar on page 4" in p["text"]
            order.append("text")
        elif p.get("type") == "image_url":
            url = p["image_url"]["url"]
            img = Image.open(io.BytesIO(base64.b64decode(url.split(",", 1)[1]))).convert("RGB")
            c["small"] &= max(img.size) <= 1600
            colors = img.getcolors(maxcolors=img.size[0] * img.size[1])
            best = None
            for count, rgb in sorted(colors, reverse=True):
                if max(rgb) - min(rgb) > 60:  # the most common saturated colour
                    best = rgb; break
            if best is None:
                best = sorted(colors, reverse=True)[1][1]  # page 6 is grey
            page = min(COLORS, key=lambda n: sum((a - b * 255) ** 2 for a, b in zip(best, COLORS[n])))
            pages.append(page); order.append("image")
        else:
            c["valid_messages"] = False
c["pages_3_to_5"] = sorted(pages) == [3, 4, 5]
c["info_interleaved"] = order[:6] == ["text", "image"] * 3
print(json.dumps(c))
""",
    "folder": r"""
import importlib.util, json, os, shutil, subprocess, sys
from pathlib import Path
py = sys.executable
work = Path.cwd()
c = {"stops_naming_all": False, "writes_nothing_on_failure": False,
     "saves_and_loads": False, "load_needs_no_documents": False}
def run(folder):
    p = subprocess.run([py, "prepare.py"], cwd=folder, capture_output=True, text=True, timeout=300)
    return p.returncode, p.stdout + p.stderr
grading = work.parent / "grading"
shutil.rmtree(grading, ignore_errors=True)
for name in ("broken", "clean"):
    shutil.copytree(work, grading / name, ignore=shutil.ignore_patterns(".venv", "docs.json", "__pycache__"))
code, out = run(grading / "broken")
c["stops_naming_all"] = code != 0 and "broken.pdf" in out and "damaged.docx" in out
c["writes_nothing_on_failure"] = code != 0 and not (grading / "broken" / "docs.json").exists()
for bad in ("broken.pdf", "damaged.docx"):
    (grading / "clean" / "docs" / bad).unlink()
code, out = run(grading / "clean")
# Outcome, whatever options the agent chose (page pictures, all rows, ...):
# every file comes back, with its text, and every picture's bytes intact.
EXPECTED = {"brief.pdf": "Project brief", "notes.txt": "Kickoff", "budget.xlsx": "servers",
            "memo.docx": "two more testers", "diagram.png": None}
import io
from PIL import Image
import attachments
os.chdir(grading / "clean")
try:
    json.loads(Path("docs.json").read_text())
    spec = importlib.util.spec_from_file_location("prepare", "prepare.py")
    m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
    def intact(loaded):
        if not isinstance(loaded, attachments.Artifacts) or loaded.errors:
            return False
        got = {Path(a["meta"]["source"]).name: a for a in loaded}
        if set(got) != set(EXPECTED) or not got["diagram.png"]["images"]:
            return False
        for name, phrase in EXPECTED.items():
            if phrase and phrase not in got[name]["text"]:
                return False
        for image in loaded.images:
            Image.open(io.BytesIO(image["bytes"])).load()
        return True
    c["saves_and_loads"] = code == 0 and intact(m.load())
    shutil.move("docs", "docs-away")
    c["load_needs_no_documents"] = c["saves_and_loads"] and intact(m.load())
except Exception as exc:
    c["error"] = repr(exc)[:300]
print(json.dumps(c))
""",
    "scans": r"""
import base64, importlib.util, io, json
from PIL import Image
spec = importlib.util.spec_from_file_location("ask", "ask.py")
m = importlib.util.module_from_spec(spec); spec.loader.exec_module(m)
body = m.request("Which brackets need action?")
raw = json.dumps(body)
c = {"valid_body": body.get("model") == "claude-sonnet-4-5" and body.get("max_tokens") == 1000,
     "question_sent": "Which brackets need action?" in raw,
     "every_page_sent": False, "request_under_32MB": len(raw.encode()) <= 32 * 2**20,
     "pictures_within_limits": True}
images = [b for msg in body.get("messages", []) if isinstance(msg.get("content"), list)
          for b in msg["content"] if isinstance(b, dict) and b.get("type") == "image"]
for b in images:
    data = base64.b64decode(b["source"]["data"])
    img = Image.open(io.BytesIO(data)); img.load()
    ok = len(data) <= 5 * 2**20 and (len(images) <= 20 or max(img.size) <= 2000)
    c["pictures_within_limits"] &= ok and len(images) <= 100
c["every_page_sent"] = len(images) == 40
c["info_request_MB"] = round(len(raw.encode()) / 2**20, 1)
print(json.dumps(c))
""",
    "install": r"""
import json, subprocess, tomllib
from pathlib import Path
c = {"uses_1_0": False, "declared": False, "prints_pages_2_3_only": False}
deps = tomllib.loads(Path("pyproject.toml").read_text()).get("project", {}).get("dependencies", [])
c["declared"] = any(d.lower().startswith("attachments") for d in deps)
env = {k: v for k, v in __import__("os").environ.items() if k != "VIRTUAL_ENV"}
p = subprocess.run(["uv", "run", "--frozen", "python", "-c",
                    "import attachments as a; print(hasattr(a, 'att'))"],
                   capture_output=True, text=True, env=env, timeout=300)
c["uses_1_0"] = p.stdout.strip().endswith("True")
p = subprocess.run(["uv", "run", "--frozen", "python", "show.py"], capture_output=True, text=True,
                   env=env, timeout=300)
out = p.stdout
c["prints_pages_2_3_only"] = (p.returncode == 0 and "PAGE-2" in out and "PAGE-3" in out
                              and "PAGE-1" not in out and "PAGE-4" not in out)
print(json.dumps(c))
""",
}


def grade(task: str, work: Path) -> dict:
    code = (
        GRADERS[task] % {"colors": repr(CHART_COLORS)}
        if task == "charts"
        else GRADERS[task]
    )
    env = {k: v for k, v in os.environ.items() if k != "VIRTUAL_ENV"}
    if task == "install":
        cmd = [sys.executable, "-c", code]
    else:
        cmd = [str(work / ".venv" / "bin" / "python"), "-c", code]
    try:
        p = subprocess.run(
            cmd, cwd=work, capture_output=True, text=True, env=env, timeout=900
        )
        checks = json.loads(p.stdout.strip().splitlines()[-1])
    except Exception as exc:
        tail = locals().get("p").stderr[-600:] if locals().get("p") else ""
        return {
            "success": False,
            "score": 0.0,
            "checks": {},
            "error": f"{exc!r} {tail}",
        }
    required = {
        k: v
        for k, v in checks.items()
        if isinstance(v, bool) and not k.startswith("info")
    }
    score = sum(required.values()) / len(required)
    return {"success": score == 1.0, "score": round(score, 2), "checks": checks}


# =============================================================================
# Expert solutions (each must score 1.0) and naive ones (must not)
# =============================================================================

SOLUTIONS = {
    "photos": {
        "classify.py": """
            from attachments import att

            def request(path: str) -> dict:
                messages = att(path).raise_for_errors().claude(
                    "Which animal is in this picture? Answer with one word.", sources=False
                )
                return {"model": "claude-sonnet-4-5", "max_tokens": 20, "messages": messages}
        """,
    },
    "charts": {
        "ask.py": """
            from attachments import att

            def messages(question: str) -> list:
                a = att("report.pdf[pages: 3-5, images: true, max_dim: 1568]").raise_for_errors()
                return a.openai(question)
        """,
    },
    "folder": {
        "prepare.py": """
            import json
            import sys

            from attachments import Artifacts, AttachmentsError, att

            def load() -> Artifacts:
                with open("docs.json") as f:
                    return Artifacts.from_wire(json.load(f))

            if __name__ == "__main__":
                try:
                    a = att("docs/").raise_for_errors()
                except AttachmentsError as e:
                    sys.exit(f"cannot read:\\n{e}")
                with open("docs.json", "w") as f:
                    json.dump(a.to_wire(), f)
        """,
    },
    "scans": {
        "ask.py": """
            from attachments import att

            def request(question: str) -> dict:
                a = att("scan.pdf[image_format: jpeg, max_dim: 1568]").raise_for_errors()
                return {"model": "claude-sonnet-4-5", "max_tokens": 1000,
                        "messages": a.claude(question, sources=False)}
        """,
    },
    "install": {
        "show.py": """
            from attachments import att

            print(att("report.pdf[pages: 2-3]").raise_for_errors().to_text(sources=False))
        """,
        "__setup__": [
            ["uv", "add", "-q", "attachments[pdf] @ {wheel}"],
        ],
    },
}

NAIVE = {
    "photos": {
        "classify.py": """
        from attachments import att

        def request(path: str) -> dict:
            return {"model": "claude-sonnet-4-5", "max_tokens": 20,
                    "messages": att(path).claude("Which animal is this?")}
    """
    },
    "charts": {
        "ask.py": """
        from attachments import att

        def messages(question: str) -> list:
            return att("report.pdf[pages: 3-5]").openai(question)
    """
    },
    "folder": {
        "prepare.py": """
        import json
        from attachments import Artifacts, att

        def load():
            return Artifacts.from_wire(json.load(open("docs.json")))

        if __name__ == "__main__":
            json.dump(att("docs/").to_wire(), open("docs.json", "w"))
    """
    },
    "scans": {
        "ask.py": """
        from attachments import att

        def request(question: str) -> dict:
            return {"model": "claude-sonnet-4-5", "max_tokens": 1000,
                    "messages": att("scan.pdf").claude(question)}
    """
    },
    "install": {
        "show.py": """
        from attachments import Attachments

        print(Attachments("report.pdf[2-3]"))
    """,
        "__setup__": [["uv", "add", "-q", "attachments"]],
    },
}


def check(wheel: Path) -> int:
    """Expert solutions must score 1.0 and naive ones less: the graders work."""
    root = CACHE / "check"
    shutil.rmtree(root, ignore_errors=True)
    failed = False
    for label, solutions, want in (
        ("expert", SOLUTIONS, True),
        ("naive", NAIVE, False),
    ):
        for task in TASKS:
            work = root / label / task / "work"
            workspace(task, "skill", work, wheel)
            env = {k: v for k, v in os.environ.items() if k != "VIRTUAL_ENV"}
            for name, body in solutions[task].items():
                if name == "__setup__":
                    for cmd in body:
                        cmd = [part.format(wheel=wheel.as_uri()) for part in cmd]
                        subprocess.run(cmd, cwd=work, check=True, env=env)
                else:
                    (work / name).write_text(textwrap.dedent(body).lstrip())
            g = grade(task, work)
            ok = g["success"] is want
            failed |= not ok
            print(
                f"{label:6} {task:8} score {g['score']:.2f} {'ok' if ok else 'WRONG'}  {g['checks'] or g.get('error')}"
            )
    return 1 if failed else 0


# =============================================================================
# Running agents
# =============================================================================


def session_metrics(session_dir: Path, arm: str) -> dict:
    m = dict(
        cost=0.0,
        turns=0,
        tool_calls=0,
        tool_errors=0,
        read_skill=False,
        contaminated=False,
        final="",
        first=None,
        last=None,
    )
    for f in sorted(session_dir.glob("*.jsonl")):
        for line in f.read_text().splitlines():
            try:
                o = json.loads(line)
            except json.JSONDecodeError:
                continue
            if o.get("type") != "message":
                continue
            m["first"] = m["first"] or o.get("timestamp")
            m["last"] = o.get("timestamp") or m["last"]
            msg = o["message"]
            if msg.get("role") == "assistant":
                m["turns"] += 1
                m["cost"] += ((msg.get("usage") or {}).get("cost") or {}).get(
                    "total", 0
                ) or 0
                for c in msg.get("content") or []:
                    if c.get("type") == "toolCall":
                        m["tool_calls"] += 1
                        args = json.dumps(c.get("arguments"))
                        if "SKILL.md" in args or "options.md" in args:
                            m["read_skill"] = True
                        if arm == "none" and FORBIDDEN.search(args):
                            m["contaminated"] = True
                        elif arm == "skill" and "Projects/attachments" in args:
                            m["contaminated"] = True
                    elif c.get("type") == "text":
                        m["final"] = c.get("text", "")
            elif msg.get("role") == "toolResult":
                text = " ".join(
                    c.get("text", "")
                    for c in msg.get("content") or []
                    if isinstance(c, dict)
                )
                if msg.get("isError") or re.search(
                    r"Traceback \(most recent|exited with code [1-9]", text
                ):
                    m["tool_errors"] += 1
    m["cost"] = round(m["cost"], 4)
    return m


def one(
    task: str, arm: str, rep: int, run_root: Path, results: Path, args, wheel: Path
) -> dict:
    base = run_root / task / f"{arm}-{rep}"
    work = base / "work"
    workspace(task, arm, work, wheel)
    env = {k: v for k, v in os.environ.items() if k != "VIRTUAL_ENV"}
    provider, model = args.model.split("/", 1)
    cmd = [
        "pi",
        "-p",
        "--mode",
        "json",
        "--no-extensions",
        "--no-skills",
        "--no-context-files",
        "--no-prompt-templates",
        "--no-themes",
        "--session-dir",
        str(base / "session"),
        "--provider",
        provider,
        "--model",
        model,
        "--thinking",
        args.thinking,
        "--append-system-prompt",
        RULES,
    ]
    if arm == "skill":
        cmd += ["--skill", str(run_root / "skill" / "attachments")]
    cmd.append(" ".join(textwrap.dedent(PROMPTS[task]).split()))
    started = time.time()
    try:
        p = subprocess.run(
            cmd, cwd=work, env=env, capture_output=True, text=True, timeout=args.timeout
        )
        status = "ok" if p.returncode == 0 else f"exit {p.returncode}"
    except subprocess.TimeoutExpired:
        status = "timeout"
    seconds = round(time.time() - started)
    metrics = session_metrics(base / "session", arm)
    graded = grade(task, work)
    result = {
        "task": task,
        "arm": arm,
        "rep": rep,
        "status": status,
        "seconds": seconds,
        "agent": metrics,
        "grade": graded,
        "work": str(work),
    }
    out = results / task / f"{arm}-{rep}"
    out.mkdir(parents=True, exist_ok=True)
    (out / "result.json").write_text(json.dumps(result, indent=1))
    for f in (base / "session").glob("*.jsonl"):
        shutil.copy(f, out / "session.jsonl")
    print(
        f"{task:8} {arm:6} #{rep}  {'✓' if graded['success'] else '✗'} {graded['score']:.2f}  "
        f"{seconds:4d}s  {metrics['tool_calls']:3d} tools ({metrics['tool_errors']} failed)"
        f"{'  read skill' if metrics['read_skill'] else ''}"
        f"{'  CONTAMINATED' if metrics['contaminated'] else ''}  {status}  "
        f"{ {k: v for k, v in graded['checks'].items() if v is False} or '' }",
        flush=True,
    )
    return result


def report(results: Path) -> str:
    rows = [json.loads(p.read_text()) for p in sorted(results.glob("*/*/result.json"))]
    lines = [
        "| task | arm | success | mean score | median time | mean tool calls | failed checks |",
        "|---|---|---|---|---|---|---|",
    ]
    for task in TASKS:
        for arm in ARMS:
            rs = [
                r
                for r in rows
                if r["task"] == task
                and r["arm"] == arm
                and not r["agent"]["contaminated"]
            ]
            if not rs:
                continue
            wins = sum(r["grade"]["success"] for r in rs)
            score = sum(r["grade"]["score"] for r in rs) / len(rs)
            secs = sorted(r["seconds"] for r in rs)[len(rs) // 2]
            tools = sum(r["agent"]["tool_calls"] for r in rs) / len(rs)
            misses: dict[str, int] = {}
            for r in rs:
                for k, v in r["grade"]["checks"].items():
                    if v is False and not k.startswith("info"):
                        misses[k] = misses.get(k, 0) + 1
            missed = ", ".join(f"{k} ({n})" for k, n in sorted(misses.items())) or "—"
            lines.append(
                f"| {task} | {arm} | {wins}/{len(rs)} | {score:.2f} | {secs}s | {tools:.0f} | {missed} |"
            )
    dirty = [
        f"{r['task']}/{r['arm']}-{r['rep']}" for r in rows if r["agent"]["contaminated"]
    ]
    if dirty:
        lines.append(f"\nLeft out (looked at forbidden files): {', '.join(dirty)}")
    return "\n".join(lines) + "\n"


def regrade(results: Path) -> int:
    """Grade every workspace of an earlier run again (after fixing a grader)."""
    for path in sorted(results.glob("*/*/result.json")):
        result = json.loads(path.read_text())
        result["grade"] = grade(result["task"], Path(result["work"]))
        path.write_text(json.dumps(result, indent=1))
    table = report(results)
    (results / "report.md").write_text(table)
    print(table)
    return 0


def build_wheel() -> Path:
    out = CACHE / "wheel"
    shutil.rmtree(out, ignore_errors=True)
    subprocess.run(
        ["uv", "build", "-q", "--wheel", "--out-dir", str(out)], cwd=REPO, check=True
    )
    return next(out.glob("*.whl"))


def main() -> int:
    ap = argparse.ArgumentParser(
        description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter
    )
    ap.add_argument("--tasks", nargs="+", default=list(TASKS), choices=TASKS)
    ap.add_argument("--arms", nargs="+", default=list(ARMS), choices=ARMS)
    ap.add_argument("--reps", type=int, default=3)
    ap.add_argument("--parallel", type=int, default=6)
    ap.add_argument(
        "--model",
        default="openai-codex/gpt-6-astra",
        help="provider/model as Pi names it",
    )
    ap.add_argument("--thinking", default="medium")
    ap.add_argument("--timeout", type=int, default=1200)
    ap.add_argument(
        "--check", action="store_true", help="grade expert and naive solutions only"
    )
    ap.add_argument("--report", type=Path, help="print the table of an earlier run")
    ap.add_argument(
        "--regrade", type=Path, help="grade an earlier run's workspaces again"
    )
    args = ap.parse_args()
    if args.report:
        print(report(args.report.resolve()))
        return 0
    if args.regrade:
        return regrade(args.regrade.resolve())
    wheel = build_wheel()
    if args.check:
        return check(wheel)
    stamp = time.strftime("%Y%m%d-%H%M%S")
    run_root, results = CACHE / stamp, HERE / "runs" / stamp
    shutil.copytree(
        SKILL,
        run_root / "skill" / "attachments",
        ignore=shutil.ignore_patterns("__pycache__"),
    )
    results.mkdir(parents=True)
    (results / "setup.json").write_text(
        json.dumps(
            {
                "model": args.model,
                "thinking": args.thinking,
                "reps": args.reps,
                "skill_sha": subprocess.run(
                    ["git", "hash-object", str(SKILL / "SKILL.md")],
                    capture_output=True,
                    text=True,
                ).stdout.strip(),
                "commit": subprocess.run(
                    ["git", "-C", str(REPO), "rev-parse", "HEAD"],
                    capture_output=True,
                    text=True,
                ).stdout.strip(),
            },
            indent=1,
        )
    )
    jobs = [
        (t, a, r)
        for r in range(1, args.reps + 1)
        for t in args.tasks
        for a in args.arms
    ]
    with futures.ThreadPoolExecutor(args.parallel) as pool:
        list(pool.map(lambda j: one(*j, run_root, results, args, wheel), jobs))
    table = report(results)
    (results / "report.md").write_text(table)
    print("\n" + table)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
