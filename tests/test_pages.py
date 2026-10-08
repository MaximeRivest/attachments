"""Page selections (``pages: 1,3,5``, ``-1``, ``-3-``, ``7-``) for PDFs and
PowerPoint, through ``att()`` and the DSL."""

from __future__ import annotations

import io
from pathlib import Path

import pytest

from attachments import att
from attachments._pages import PageSelectionError, parse_pages
from attachments.deps import check_dep


class TestParse:
    @pytest.mark.parametrize(
        ("text", "total", "pages"),
        [
            ("3", 5, [3]),
            ("2-4", 5, [2, 3, 4]),
            ("4-", 5, [4, 5]),
            ("1,3,5", 5, [1, 3, 5]),
            ("-1", 5, [5]),
            ("-2", 5, [4]),
            ("-3-", 5, [3, 4, 5]),
            ("2--2", 5, [2, 3, 4]),
            ("5,1,5", 5, [1, 5]),  # document order, once each
            ("1-3,2-4", 5, [1, 2, 3, 4]),
            ("4-9", 5, [4, 5]),  # clipped to the document
            ("-9", 5, []),
        ],
    )
    def test_selections(self, text, total, pages):
        assert [i + 1 for i in parse_pages(text).indices(total)] == pages

    @pytest.mark.parametrize("bad", ["0", "x", "5-2", "1,,x", "", True, 2.5, [1, 2, 3]])
    def test_invalid(self, bad):
        with pytest.raises(PageSelectionError):
            parse_pages(bad)

    def test_two_number_list_is_a_range(self):
        # JSON carries the DSL range (2, 5) as [2, 5].
        assert parse_pages([2, 5]).indices(9) == [1, 2, 3, 4]


@pytest.fixture
def five_page_pdf(tmp_path: Path) -> Path:
    pymupdf = pytest.importorskip("pymupdf")
    doc = pymupdf.open()
    for i in range(5):
        doc.new_page().insert_text((72, 72), f"Page number {i + 1}")
    path = tmp_path / "five.pdf"
    doc.save(path)
    return path


def pages_of(artifact) -> list[int]:
    return [s["page"] for s in artifact["meta"]["segments"]]


@pytest.mark.skipif(not check_dep("pdf").available, reason="pdf extra needed")
class TestPdf:
    @pytest.mark.parametrize(
        ("dsl", "pages"),
        [
            ("pages: 1,3,5", [1, 3, 5]),  # unquoted: the DSL continues the value
            ('pages: "1,3,5"', [1, 3, 5]),
            ("pages: -1", [5]),
            ("pages: -2-", [4, 5]),
            ("pages: 4-", [4, 5]),
            ("page: 2", [2]),
        ],
    )
    def test_dsl(self, five_page_pdf, dsl, pages):
        [a] = att(f"{five_page_pdf}[{dsl}]")
        assert pages_of(a) == pages
        assert "warnings" not in a["meta"]
        assert all(f"Page number {p}" in a["text"] for p in pages)

    def test_keyword_and_other_options(self, five_page_pdf):
        [a] = att(str(five_page_pdf), pages="1,-1", images=True, max_dim=200)
        assert pages_of(a) == [1, 5]
        assert [im["page"] for im in a["images"]] == [1, 5]
        assert a["meta"]["extra"]["page_selection"] == "1,-1"

    def test_max_pages_caps_a_selection(self, five_page_pdf):
        [a] = att(str(five_page_pdf), pages="2-", max_pages=2)
        assert pages_of(a) == [2, 3]

    def test_selection_outside_the_document_is_reported(self, five_page_pdf):
        for spec in ("7-9", "9", "-9"):
            [a] = att(str(five_page_pdf), pages=spec)
            assert a["text"] == ""
            assert "selects no page of this 5-page document" in a["meta"]["warnings"][0]
            assert "ocr_hint" not in a["meta"]["extra"]  # not "a scan"

    def test_invalid_value_warns_and_reads_everything(self, five_page_pdf):
        [a] = att(str(five_page_pdf), pages="1-x")
        assert pages_of(a) == [1, 2, 3, 4, 5]
        assert "Invalid value for option 'pages'" in a["meta"]["warnings"][0]

    def test_pdf2image_fallback_draws_each_run(self, five_page_pdf, monkeypatch):
        import sys

        from PIL import Image

        from attachments._processors import pdf as pdf_module

        calls: list[tuple[int, int]] = []

        class FakePdf2image:
            @staticmethod
            def pdfinfo_from_bytes(data):
                return {"Pages": 5}

            @staticmethod
            def convert_from_bytes(data, dpi, first_page, last_page, fmt):
                calls.append((first_page, last_page))
                return [
                    Image.new("RGB", (10, 10)) for _ in range(first_page, last_page + 1)
                ]

        monkeypatch.setattr(
            pdf_module,
            "_render_pages_with_pymupdf",
            lambda *a, **k: ([], None, {}),
        )
        monkeypatch.setitem(sys.modules, "pdf2image", FakePdf2image)
        [a] = att(str(five_page_pdf), pages="1-2,5", images=True)
        assert calls == [(1, 2), (5, 5)]
        assert [im["page"] for im in a["images"]] == [1, 2, 5]


@pytest.mark.skipif(not check_dep("pptx").available, reason="pptx extra needed")
class TestPptx:
    @pytest.fixture
    def deck(self, tmp_path: Path) -> Path:
        import pptx

        prs = pptx.Presentation()
        for title in ("One", "Two", "Three", "Four"):
            prs.slides.add_slide(prs.slide_layouts[1]).shapes.title.text = title
        path = tmp_path / "deck.pptx"
        buf = io.BytesIO()
        prs.save(buf)
        path.write_bytes(buf.getvalue())
        return path

    @pytest.mark.parametrize(
        ("dsl", "slides"),
        [("pages: 2-3", [2, 3]), ("slides: 1,4", [1, 4]), ("pages: -1", [4])],
    )
    def test_slides(self, deck, dsl, slides):
        [a] = att(f"{deck}[{dsl}]")
        assert pages_of(a) == slides
        assert a["meta"]["extra"]["slides_selected"] == slides
        assert a["meta"]["extra"]["slide_count"] == 4

    def test_no_slide_selected(self, deck):
        [a] = att(str(deck), pages="7")
        assert a["meta"]["warnings"] == [
            "pages: 7 selects no slide of this 4-slide deck"
        ]


def test_select_with_commas_needs_no_quotes(tmp_path: Path):
    pytest.importorskip("bs4")
    page = tmp_path / "p.html"
    page.write_text("<h1>A</h1><p>b</p><h2>c</h2>")
    [a] = att(f"{page}[select: h1, h2]")
    assert a["text"] == "# A\n\n## c"
