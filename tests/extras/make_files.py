"""Write the sample files of tests/extras/check.py (run with the dev environment).

    uv run python tests/extras/make_files.py

One small file per format, made here so they can be remade; check.py only
reads them, in environments that may lack every library used below.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

HERE = Path(__file__).parent / "files"


def main() -> None:
    import openpyxl
    import pymupdf
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    from pptx import Presentation

    HERE.mkdir(exist_ok=True)
    font = ImageFont.load_default(size=48)

    doc = pymupdf.open()
    doc.new_page(width=300, height=150).insert_text((20, 80), "Text layer", fontsize=24)
    doc.save(HERE / "text.pdf", garbage=4, deflate=True)

    # A scan: a picture of a page, no text layer.
    picture = Image.new("L", (600, 300), 255)
    ImageDraw.Draw(picture).text((30, 120), "SCANNED 42", fill=0, font=font)
    buf = io.BytesIO()
    picture.save(buf, format="PNG", optimize=True)
    scan = pymupdf.open()
    page = scan.new_page(width=300, height=150)
    page.insert_image(page.rect, stream=buf.getvalue())
    scan.save(HERE / "scan.pdf", garbage=4, deflate=True)

    sign = Image.new("RGB", (320, 100), "white")
    ImageDraw.Draw(sign).text((10, 25), "HELLO", fill="black", font=font)
    sign.save(HERE / "sign.png", optimize=True)

    d = Document()
    d.add_heading("Report", 1)
    d.add_paragraph("A Word paragraph.")
    d.save(HERE / "report.docx")

    p = Presentation()
    slide = p.slides.add_slide(p.slide_layouts[1])
    slide.shapes.title.text = "Slide title"
    slide.placeholders[1].text = "A bullet"
    p.save(HERE / "slides.pptx")

    wb = openpyxl.Workbook()
    wb.active.title = "Sales"
    wb.active.append(["region", "total"])
    wb.active.append(["north", 12])
    wb.save(HERE / "book.xlsx")

    (HERE / "page.html").write_text(
        "<html><head><title>Page</title></head><body><h1>Heading</h1>"
        "<p>An HTML paragraph.</p></body></html>\n"
    )
    (HERE / "drawing.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="60">'
        '<text x="10" y="40" font-size="30">Hi SVG</text></svg>\n'
    )
    (HERE / "data.csv").write_text("a,b\n1,2\n3,4\n")
    (HERE / "notes.md").write_text("# Notes\n\nPlain text needs no extra.\n")
    notebook = {
        "cells": [
            {"cell_type": "markdown", "metadata": {}, "source": ["# A notebook"]}
        ],
        "metadata": {},
        "nbformat": 4,
        "nbformat_minor": 5,
    }
    (HERE / "notes.ipynb").write_text(json.dumps(notebook) + "\n")
    for f in sorted(HERE.iterdir()):
        print(f"{f.stat().st_size:>7}  {f.name}")


if __name__ == "__main__":
    main()
