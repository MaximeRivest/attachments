"""Tests for to_parts / Artifacts.parts and the presenters built on it.

Covers:
    - Text-only input: one text part equal to render_text (both sources modes)
    - Interleaving: each page's text, then that page's images (real PDF,
      real PPTX, hand-built edge cases: gaps, unpaged images, bad segments)
    - sources=False: no file name anywhere in the output
    - interleave=False: the text-first layout (render_text, then images)
    - prompt: its own final part; empty prompt adds nothing
    - Claude/OpenAI output: the same parts in each provider's wire format
    - Wire-form artifacts (bytes_b64) give the same parts as in-process ones
"""

from __future__ import annotations

import base64
import io
import json

import pytest

from attachments import Artifacts, att, render_text, to_parts
from attachments.deps import check_dep
from attachments.render import to_claude_content, to_openai_messages
from attachments.types import make_artifact


def _img(name: str, page: int | None = None, payload: bytes = b"\x89PNG") -> dict:
    item = {"name": name, "mimetype": "image/png", "bytes": payload}
    if page is not None:
        item["page"] = page
    return item


def _paged(text_pages: list[str], images: list[dict], source: str = "doc.pdf"):
    """An artifact whose text is pages joined by blank lines, with segments."""
    segments, offset = [], 0
    for number, page_text in enumerate(text_pages, start=1):
        segments.append(
            {
                "kind": "page",
                "label": f"page {number}",
                "start": offset,
                "end": offset + len(page_text),
                "page": number,
            }
        )
        offset += len(page_text) + 2
    return make_artifact(
        text="\n\n".join(text_pages),
        images=images,
        meta={"source": source, "segments": segments},
    )


def _shape(parts: list[dict]) -> list[str]:
    """Text parts as their text, image parts as 'img:<payload>'."""
    return [
        p["text"]
        if p["type"] == "text"
        else "img:" + base64.b64decode(p["data"]).decode("latin-1")
        for p in parts
    ]


# =============================================================================
# Text-only input
# =============================================================================


@pytest.mark.parametrize("sources", [True, False])
def test_text_only_input_is_one_part_equal_to_render_text(sources: bool):
    arts = [
        make_artifact(text="  Alpha.  ", meta={"source": "a.txt"}),
        make_artifact(text="", meta={"source": "empty.txt"}),
        make_artifact(text="Beta.", meta={"source": "b.txt"}),
    ]
    expected = render_text(arts, sources=sources)
    assert to_parts(arts, sources=sources) == [{"type": "text", "text": expected}]


def test_empty_input_gives_no_parts():
    assert to_parts([]) == []
    assert to_parts([make_artifact(meta={"source": "x"})]) == []


# =============================================================================
# Interleaving (hand-built)
# =============================================================================


def test_each_page_text_is_followed_by_its_image():
    art = _paged(
        ["one", "two", "three"],
        [_img("p1", 1, b"1"), _img("p2", 2, b"2"), _img("p3", 3, b"3")],
    )
    assert _shape(to_parts([art], sources=False)) == [
        "one",
        "img:1",
        "two",
        "img:2",
        "three",
        "img:3",
    ]


def test_pages_without_images_keep_their_text_together():
    art = _paged(["one", "two", "three"], [_img("p3", 3, b"3")])
    assert _shape(to_parts([art], sources=False)) == ["one\n\ntwo\n\nthree", "img:3"]


def test_several_images_on_one_page_keep_their_order():
    art = _paged(["one", "two"], [_img("a", 1, b"a"), _img("b", 1, b"b")])
    assert _shape(to_parts([art], sources=False)) == ["one", "img:a", "img:b", "two"]


def test_image_for_a_page_without_segment_goes_before_the_next_page():
    art = _paged(["one", "two", "three"], [_img("x", 2, b"x")])
    # Drop the page-2 segment: the image lands right before page 3's text
    # (after the unsegmented text in between, which is page 2's).
    art["meta"]["segments"].pop(1)
    assert _shape(to_parts([art], sources=False)) == [
        "one\n\ntwo",
        "img:x",
        "three",
    ]


def test_unpaged_and_out_of_range_images_go_at_the_end():
    art = _paged(
        ["one", "two"],
        [_img("u", None, b"u"), _img("p1", 1, b"1"), _img("p9", 9, b"9")],
    )
    assert _shape(to_parts([art], sources=False)) == [
        "one",
        "img:1",
        "two",
        "img:u",
        "img:9",
    ]


def test_segments_without_page_numbers_fall_back_to_text_then_images():
    art = _paged(["one", "two"], [_img("p1", 1, b"1")])
    for segment in art["meta"]["segments"]:
        del segment["page"]
    assert _shape(to_parts([art], sources=False)) == ["one\n\ntwo", "img:1"]


def test_labels_are_never_parsed_for_page_numbers():
    art = _paged(["one", "two"], [_img("p2", 2, b"2")])
    for segment in art["meta"]["segments"]:
        del segment["page"]  # only the labels "page 1"/"page 2" remain
    assert _shape(to_parts([art], sources=False)) == ["one\n\ntwo", "img:2"]


def test_malformed_segments_are_ignored_not_fatal():
    art = _paged(["one", "two"], [_img("p1", 1, b"1")])
    art["meta"]["segments"].append({"kind": "page", "page": "3", "start": None})
    art["meta"]["segments"].append({"kind": "page", "page": True, "start": 0, "end": 1})
    assert _shape(to_parts([art], sources=False)) == ["one", "img:1", "two"]


def test_images_without_payload_are_skipped():
    art = _paged(
        ["one", "two"], [{"name": "ghost.png", "mimetype": "image/png", "page": 1}]
    )
    assert _shape(to_parts([art], sources=False)) == ["one\n\ntwo"]


# =============================================================================
# sources
# =============================================================================


def test_sources_true_puts_the_header_on_the_first_part():
    art = _paged(["one", "two"], [_img("p1", 1, b"1")], source="report.pdf")
    assert _shape(to_parts([art])) == ["## report.pdf\none", "img:1", "two"]


def test_image_only_artifact_gets_a_header_part_then_the_image():
    photo = make_artifact(
        images=[_img("tabby_cat.png", payload=b"c")], meta={"source": "tabby_cat.png"}
    )
    assert _shape(to_parts([photo])) == ["## tabby_cat.png", "img:c"]
    assert _shape(to_parts([photo], sources=False)) == ["img:c"]


def test_sources_false_leaks_no_name_anywhere():
    photo = make_artifact(
        images=[_img("tabby_cat.png", payload=b"c")], meta={"source": "tabby_cat.png"}
    )
    doc = _paged(["one"], [_img("secret-page-1.png", 1, b"1")], source="secret.pdf")
    for interleave in (True, False):
        dumped = json.dumps(
            to_parts([photo, doc], sources=False, interleave=interleave)
        )
        assert "tabby" not in dumped
        assert "secret" not in dumped
        assert "## " not in dumped


def test_artifacts_merge_into_one_text_part_across_files():
    first = make_artifact(text="A", meta={"source": "a.txt"})
    photo = make_artifact(
        images=[_img("p.png", payload=b"p")], meta={"source": "p.png"}
    )
    last = make_artifact(text="B", meta={"source": "b.txt"})
    assert _shape(to_parts([first, photo, last])) == [
        "## a.txt\nA\n\n## p.png",
        "img:p",
        "## b.txt\nB",
    ]


# =============================================================================
# interleave=False and prompt
# =============================================================================


def test_interleave_false_is_render_text_then_every_image():
    art = _paged(["one", "two"], [_img("p1", 1, b"1"), _img("p2", 2, b"2")])
    photo = make_artifact(
        images=[_img("c.png", payload=b"c")], meta={"source": "c.png"}
    )
    parts = to_parts([art, photo], interleave=False)
    assert _shape(parts) == [render_text([art, photo]), "img:1", "img:2", "img:c"]


def test_prompt_is_its_own_last_part_and_empty_prompt_adds_nothing():
    art = make_artifact(text="Body.", meta={"source": "a.txt"})
    assert _shape(to_parts([art], prompt="Summarize.")) == [
        "## a.txt\nBody.",
        "Summarize.",
    ]
    assert to_parts([art], prompt="") == to_parts([art])
    assert to_parts([], prompt="Hi") == [{"type": "text", "text": "Hi"}]


# =============================================================================
# Presenters are built from parts
# =============================================================================


def test_claude_and_openai_follow_the_parts_order_exactly():
    art = _paged(["one", "two"], [_img("p1", 1, b"1"), _img("p2", 2, b"2")])
    parts = to_parts([art], prompt="Go", sources=False)
    claude = to_claude_content([art], prompt="Go", sources=False)
    openai = to_openai_messages([art], prompt="Go", sources=False)[0]["content"]
    assert [b["type"] for b in claude] == ["text", "image", "text", "image", "text"]
    assert [p["type"] for p in openai] == [
        "text",
        "image_url",
        "text",
        "image_url",
        "text",
    ]
    for part, block, oa in zip(parts, claude, openai, strict=True):
        if part["type"] == "text":
            assert block == oa == {"type": "text", "text": part["text"]}
        else:
            assert block["source"] == {
                "type": "base64",
                "media_type": part["media_type"],
                "data": part["data"],
            }
            assert oa["image_url"]["url"] == (
                f"data:{part['media_type']};base64,{part['data']}"
            )


def test_artifacts_shortcuts_pass_options_through():
    a = Artifacts([_paged(["one", "two"], [_img("p1", 1, b"1")], source="r.pdf")])
    assert a.parts(sources=False, prompt="Q") == to_parts(a, sources=False, prompt="Q")
    assert a.claude("Q", sources=False)[0]["content"][0] == {
        "type": "text",
        "text": "one",
    }
    assert a.openai("Q", interleave=False)[0]["content"][0]["text"].startswith(
        "## r.pdf"
    )


def test_wire_form_gives_the_same_parts():
    a = Artifacts([_paged(["one", "two"], [_img("p1", 1, b"1")])])
    assert to_parts(a.to_wire()) == a.parts()


# =============================================================================
# Real documents
# =============================================================================


@pytest.mark.skipif(
    not check_dep("pdf-images").available, reason="PyMuPDF needed to build PDFs"
)
def test_real_pdf_interleaves_page_text_and_page_render(tmp_path):
    pymupdf = pytest.importorskip("pymupdf")
    doc = pymupdf.open()
    for word in ("alpha", "beta", "gamma"):
        doc.new_page().insert_text((72, 72), f"{word} words")
    path = tmp_path / "three.pdf"
    doc.save(path)
    doc.close()

    a = att(f"{path}[images: true, dpi: 36]")
    parts = a.parts()
    assert [p["type"] for p in parts] == ["text", "image"] * 3
    assert parts[0]["text"] == "## three.pdf\nalpha words"
    assert parts[2]["text"] == "beta words"
    assert parts[4]["text"] == "gamma words"
    renders = [image["bytes"] for image in a.images]
    assert [base64.b64decode(p["data"]) for p in parts[1::2]] == renders

    # A page range keeps text and pictures aligned (pages 2-3 only).
    sub = att(f"{path}[pages: 2-3, images: true, dpi: 36]")
    sub_parts = sub.parts(sources=False)
    assert [p.get("text") for p in sub_parts[::2]] == ["beta words", "gamma words"]
    assert [i["page"] for i in sub.images] == [2, 3]


@pytest.mark.skipif(not check_dep("pptx").available, reason="python-pptx needed")
def test_real_pptx_puts_slide_picture_after_its_slide():
    from pptx import Presentation
    from pptx.util import Inches

    png = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (4, 4), "red").save(png, format="PNG")
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[5])
    first.shapes.title.text = "Cover"
    second = prs.slides.add_slide(prs.slide_layouts[5])
    second.shapes.title.text = "Chart"
    second.shapes.add_picture(io.BytesIO(png.getvalue()), Inches(1), Inches(1))
    third = prs.slides.add_slide(prs.slide_layouts[5])
    third.shapes.title.text = "Thanks"
    buf = io.BytesIO()
    prs.save(buf)

    from attachments._processors import processors

    art = processors[".pptx"](
        buf.getvalue(), filename="deck.pptx", embedded_images=True
    )
    art["meta"]["source"] = "deck.pptx"
    assert [s["page"] for s in art["meta"]["segments"]] == [1, 2, 3]
    parts = to_parts([art], sources=False)
    assert [p["type"] for p in parts] == ["text", "image", "text"]
    assert "Chart" in parts[0]["text"] and "Thanks" not in parts[0]["text"]
    assert "Thanks" in parts[2]["text"]
